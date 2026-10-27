#!/usr/bin/env python3
"""Every page and every Settings section in a real browser, with demo data: no page error, no broken request, and the Cards Excel export opens.
  python3 tests/browser/pages_test.py build/site-update-X.Y.Z.zip
Needs PHP 8.3+ and Playwright with Chromium (or --browser firefox|webkit); openpyxl is used to read the Excel file when it is installed.
Nothing leaves this computer: every request to another site is blocked."""
import io, json, os, re, shutil, subprocess, sys, tempfile, time, urllib.request, http.cookiejar, zipfile
from playwright.sync_api import sync_playwright

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'lib'))
from helpers import ROOT, FIX, SEED
from helpers import browser_name, launch, browser_login, CSP_REPORTER, TRACK_REQUESTS, settle, until
BROWSER = browser_name()
if len(sys.argv) < 2: sys.exit(__doc__)
ZIP = sys.argv[1]; PORT = 8195; B = 'http://127.0.0.1:%d' % PORT; UA = 'PagesTest/1.0'; PW = 'correct horse battery staple'
D = tempfile.mkdtemp(); zipfile.ZipFile(ZIP).extractall(D)
php = subprocess.Popen(['php', '-S', '127.0.0.1:%d' % PORT, '-t', D], env=dict(os.environ, PHP_CLI_SERVER_WORKERS='6'), stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
time.sleep(1)
res = []
def check(name, ok, extra=''):
    res.append(bool(ok)); print(('PASS ' if ok else 'FAIL ') + name + ('' if ok or not extra else '  -> ' + str(extra)[:300]))

cj = http.cookiejar.CookieJar(); op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj))
def call(a, b=None):
    rq = urllib.request.Request(B + '/web/hub/core/api.php?action=' + a, data=json.dumps(b).encode() if b is not None else None, headers={'User-Agent': UA, 'Content-Type': 'application/json'}, method='POST' if b is not None else 'GET')
    return json.loads(op.open(rq).read())
try:
    call('state')
    code = re.search(r'([A-Z0-9]{4}(?:-[A-Z0-9]{4})+|[A-Z0-9]{8,})', open(D + '/site-data/core/SETUP-CODE.php').read().replace('<?php', '')).group(1)
    check('P1 a new site can be set up', call('setup', {'code': code, 'password': PW}).get('ok'))
    subprocess.run(['php', SEED, D], capture_output=True)
    ver = re.search(r"define\('HUB', '([^']+)'\)", open(D + '/web/hub/core/lib/boot.php').read()).group(1)
    cookies = [{'name': c.name, 'value': c.value, 'url': B + '/'} for c in cj]
    panes = re.findall(r"\['([a-z]+)', '[A-Z][^']*', ICON", open(D + '/web/hub/docs/src/hub/core/shell-admin.src.js', encoding='utf-8').read())
    with sync_playwright() as p:
        b = launch(p, BROWSER); print('browser: %s %s' % (BROWSER, b.version))
        ctx = b.new_context(user_agent=UA, accept_downloads=True, viewport={'width': 1280, 'height': 860})
        ctx.add_init_script("try{localStorage.setItem('ix-seen-ver','%s')}catch(e){}" % ver)
        ctx.add_init_script(CSP_REPORTER)
        ctx.add_init_script(TRACK_REQUESTS)
        # PHP's test server ignores .htaccess, so each page gets the security policy Apache would send, read from its own .htaccess
        CSP = {}
        for app in ('links', 'info', 'cards', 'tasks', 'tools', 'settings'):
            m = re.search(r'Header (?:always )?set Content-Security-Policy "([^"]+)"', open(D + '/web/%s/.htaccess' % app).read()) if os.path.isfile(D + '/web/%s/.htaccess' % app) else None
            if m: CSP['/web/%s/' % app] = m.group(1)
        def handle(r):
            u = r.request.url
            if not (u.startswith(B) or u.startswith(('data:', 'blob:'))): return r.abort()
            path = u[len(B):].split('#')[0].split('?')[0]
            # WebKit fails the requests of a page that Playwright fetched and handed over (seen in CI: P6 failed, P7 passed),
            # so there each page keeps only its own <meta> policy; Chromium, Firefox and the Apache job check the headers.
            if BROWSER != 'webkit' and r.request.resource_type == 'document' and path in CSP:
                resp = r.fetch(); h = dict(resp.headers); h['content-security-policy'] = CSP[path]; return r.fulfill(response=resp, headers=h)
            r.continue_()
        ctx.route('**/*', handle)
        pg = ctx.new_page(); errs = []; bad = []
        why = browser_login(pg, B, PW)
        check('P0 the browser logs in through the login dialog', why == '', why)
        pg.on('pageerror', lambda e: errs.append(str(e)[:200]))
        refused = []; tools_main = set()
        pg.on('console', lambda m: refused.append(m.text[:160]) if m.type in ('error', 'warning') and ('Content Security Policy' in m.text or 'Content-Security-Policy' in m.text or 'Refused to' in m.text or 'React error' in m.text) else None)
        pg.on('request', lambda r: tools_main.add(r.url.replace(B, '')) if '/tools/assets/app-' in r.url and r.url.split('?')[0].endswith('.js') else None)
        pg.on('response', lambda r: bad.append((r.status, r.url.replace(B, ''))) if r.url.startswith(B) and r.status >= 400 else None)
        problems = []; noicon = []
        for name, url in [('Home', '/'), ('Links', '/web/links/'), ('Info', '/web/info/'), ('Cards', '/web/cards/'), ('Tasks', '/web/tasks/'), ('Tools', '/web/tools/'), ('Tools: a tool', '/web/tools/#/tool/image-resize')] + [('Settings ' + s, '/web/settings/#' + s) for s in panes]:
            errs.clear(); bad.clear()
            pg.goto(B + url, wait_until='domcontentloaded'); settle(pg)   # every request of the page has answered
            if errs or bad: problems.append((name, errs[:1], bad[:2]))
            ic = pg.evaluate("(document.querySelector('link[rel~=icon]') || {}).href || ''")
            if not ic.endswith(('/brand/favicon-64.png', '/brand/logo-192.png')) and '/media/' not in ic: noicon.append((name, ic[:60]))
        check('P2 every page and all %d Settings sections load with no page error and no broken request' % len(panes), len(panes) >= 8 and not problems, problems)
        check('P9 every page, Tools included, shows the site logo as its tab icon', not noicon, noicon)
        check('P4 with the real security policies, nothing is refused and React reports no error', len(CSP) >= 5 and not refused, (len(CSP), refused[:3]))
        check('P5 Tools loads its main file once', len(tools_main) == 1, sorted(tools_main))
        pg.goto(B + '/web/cards/'); until(lambda: pg.evaluate("!!(window.HUB && HUB.isAdmin && HUB.isAdmin())"), 10000)
        admin = pg.evaluate("!!(window.HUB && HUB.isAdmin && HUB.isAdmin())")
        diag = '' if admin else (
            [(c['name'], c['domain'], c['path'], c.get('sameSite'), c.get('secure'), c.get('expires')) for c in ctx.cookies()],
            pg.evaluate("fetch('../hub/core/api.php?action=state', {credentials: 'same-origin'}).then(r => r.json()).then(j => ({admin: j.admin}), e => String(e))"))
        check('P6 the browser is logged in (the login cookie was accepted)', admin, diag)
        # P7: the same, in a context without the test's request routing (the routing hands each page to the browser with its
        # security policy added). If P7 passes where P6 fails, the difference is in the routing, not in the site.
        c7 = b.new_context(user_agent=UA, viewport={'width': 1280, 'height': 860}); c7.add_init_script(TRACK_REQUESTS); p7 = c7.new_page()
        why7 = browser_login(p7, B, PW); p7.goto(B + '/web/cards/'); until(lambda: p7.evaluate("!!(window.HUB && HUB.isAdmin && HUB.isAdmin())"), 10000)
        check('P7 logged in on a second page without the test routing', p7.evaluate("!!(window.HUB && HUB.isAdmin && HUB.isAdmin())"), why7)
        c7.close()
        # P8: Text Sizes now opens from each page's own settings dialog, and is no longer its own item in the pencil menu
        p8 = []
        for app, dlg, btn, done in (('links', 'setDlg', 'sTextBtn', "document.getElementById('textDlg').open"),
                                    ('info', 'setdlg', 'setsizes', "!!document.querySelector('#ix-dlg.open #ix-tzs')"),
                                    ('cards', 'tdlg', 'tsizes', "!!document.querySelector('#ix-dlg.open #ix-tzs')")):
            pg.goto(B + '/web/%s/' % app, wait_until='domcontentloaded'); until(lambda: pg.evaluate("document.querySelectorAll('[data-hub=item]').length > 0"), 10000); settle(pg)
            item = pg.evaluate("document.querySelectorAll('[data-hub=item][data-hub-label=\"Text Sizes\"]').length")
            pg.evaluate("document.getElementById('%s').showModal(); document.getElementById('%s').click()" % (dlg, btn))
            until(lambda: pg.evaluate(done) and not pg.evaluate("document.getElementById('%s').open" % dlg), 5000)
            ok = pg.evaluate(done) and not pg.evaluate("document.getElementById('%s').open" % dlg)
            if item or not ok: p8.append((app, item, ok))
        check('P8 Text Sizes opens from the Settings dialog of Links, Info and Cards (and is not a separate pencil-menu item)', not p8, p8)
        data = b''
        try:
            with pg.expect_download(timeout=20000) as dl: pg.evaluate("document.querySelector('[data-a=xlsx]').click()")
            data = open(dl.value.path(), 'rb').read()
        except Exception as e:
            check('P3 Cards exports a file (no download came)', False, str(e)[:120])
        try:
            import openpyxl
            if not data: raise ValueError('no file')
            wb = openpyxl.load_workbook(io.BytesIO(data)); ws = wb['Cards']; hdr = [c.value for c in ws[1]]
            ok = 'Limit' in hdr and isinstance(ws.cell(2, hdr.index('Limit') + 1).value, (int, float)) and '₹' in ws.cell(2, hdr.index('Limit') + 1).number_format
            detail = (wb.sheetnames, hdr[:5])
        except ValueError:
            ok = False; detail = 'no file'
        except ImportError:
            z = zipfile.ZipFile(io.BytesIO(data)); ok = 'xl/worksheets/sheet1.xml' in z.namelist(); detail = z.namelist()
        check('P3 Cards exports a real Excel file with numbers and the money format', data[:2] == b'PK' and ok, detail)
        # P10-P16: the 1.48.5 Cards page (Table or By Holder sections, four total boxes, holder chips) and the tidied pencil menus
        def items(app):
            pg.goto(B + '/web/%s/' % app, wait_until='domcontentloaded'); until(lambda: pg.evaluate("document.querySelectorAll('[data-hub=item]').length > 0"), 10000); settle(pg)
            return pg.evaluate("[...document.querySelectorAll('[data-hub=item]')].map((e, i) => [+(e.dataset.hubOrder || 0), i, e.dataset.hubLabel]).sort((x, y) => x[0] - y[0] || x[1] - y[1]).map(x => x[2])")
        want = {'links': ['Colors', 'Fetch Icons', 'Import Bookmarks', 'Links Settings'], 'info': ['Export', 'Info Settings'],
                'tasks': ['Add Task', 'Spreadsheet (CSV)', 'Calendar File (.ics)', 'Print or PDF', 'Tasks Settings'],
                'cards': ['Add Card', 'Holders', 'Banks', 'PDF (print)', 'Excel Spreadsheet', 'Cards Settings']}
        got = {k: items(k) for k in want}
        check('P10 the pencil menus: Links (Colors, Fetch Icons, Import Bookmarks, Links Settings), Info, Tasks and Cards name their own settings', got == want, got)
        pg.goto(B + '/web/cards/', wait_until='domcontentloaded'); pg.evaluate("localStorage.removeItem('cards-view'); localStorage.removeItem('cards-fold')")
        pg.reload(wait_until='domcontentloaded'); pg.wait_for_selector('#list .row.item', timeout=15000)
        grp = lambda: pg.evaluate("[...document.querySelectorAll('#list .grp .gn')].map(e => e.textContent)")
        rows = lambda: pg.evaluate("document.querySelectorAll('#list .row.item').length")
        stat = lambda: pg.evaluate("[...document.querySelectorAll('#stats .stat')].map(e => e.querySelector('.sl').textContent + '=' + e.querySelector('b').textContent)")
        first = (len(grp()), rows(), pg.evaluate("getComputedStyle(document.querySelector('#list .row.item .co')).display"), stat())
        check('P11 Cards opens as the Table (every card, with a Holder column) under four boxes: Cards, Lifetime Free, Charges a Year, Expiring Soon; the first chip is All Cards',
              first[:3] == (0, 9, 'block') and first[3] == ['Cards=9', 'Lifetime Free=4 of 9', 'Charges a Year=₹8,499', 'Expiring Soon=1'] and pg.evaluate("document.querySelector('#fchips [data-h=All]').textContent") == 'All Cards'
              and pg.evaluate("document.querySelector('#fview [aria-pressed=true]').textContent") == 'Table', first)
        pg.click('#fview [data-v=holder]'); until(lambda: len(grp()) > 0, 5000)
        check('P12 By Holder: one section per holder with its totals, and the choice is kept',
              grp() == ['Alex', 'Sam', 'Jordan'] and rows() == 9 and pg.evaluate("document.querySelector('#list .grp .gh').textContent.includes('4 cards')") and pg.evaluate("localStorage.getItem('cards-view')") == 'holder', grp())
        pg.click('#list .gh[data-fold="Jordan"]'); until(lambda: rows() < 9, 5000); folded = rows()
        pg.reload(wait_until='domcontentloaded'); pg.wait_for_selector('#list .row.item', timeout=15000)
        check('P13 a holder section folds, and stays folded (and By Holder stays on) after a reload', folded == 7 and rows() == 7 and len(grp()) == 3, (folded, rows()))
        pg.click('#list .gh[data-fold="Jordan"]'); pg.click('#fview [data-v=list]'); until(lambda: not grp() and rows() == 9, 5000)
        pg.click('#fchips [data-h="Sam"]'); until(lambda: rows() < 9, 5000)
        sam = (rows(), stat()[0], pg.evaluate("!document.getElementById('fclear').hidden"))
        pg.click('#fclear'); until(lambda: rows() == 9, 5000)
        check('P16 a holder chip filters to that holder (the boxes follow), and Clear brings every card back', sam == (3, 'Cards=3', True) and rows() == 9, sam)
        pg.click('#list .row.item[data-id] .cn small'); until(lambda: pg.evaluate("document.getElementById('dlg').open"), 5000)
        opened = pg.evaluate("document.getElementById('dlg').open"); pg.evaluate("document.getElementById('dlg').close()")
        pg.evaluate("document.getElementById('tdlg').showModal()"); pg.click('#tcols'); until(lambda: pg.evaluate("document.getElementById('odlg').open && document.querySelectorAll('#olist .orow').length > 0"), 5000)
        cols = pg.evaluate("document.getElementById('odlg').open && document.querySelectorAll('#olist .orow').length"); pg.evaluate("document.getElementById('odlg').close()")
        check('P14 clicking a card row opens Edit Card, and Cards Settings opens Column Order', opened and cols == 11, (opened, cols))
        pctx = b.new_context(user_agent=UA, viewport={'width': 390, 'height': 844}); pctx.add_cookies(cookies); pctx.add_init_script("try{localStorage.setItem('ix-seen-ver','%s')}catch(e){}" % ver)
        pctx.route('**/*', handle); pp = pctx.new_page(); pp.goto(B + '/web/cards/', wait_until='domcontentloaded'); pp.wait_for_selector('#list .row.item', timeout=15000)
        ph = pp.evaluate("({sw: document.documentElement.scrollWidth - innerWidth, head: getComputedStyle(document.querySelector('#list .row.head')).display, tall: Math.max(...[...document.querySelectorAll('#list .row.item')].map(r => r.getBoundingClientRect().height))})")
        check('P15 Cards on a phone: nothing scrolls sideways, no column headings, and each card takes two short lines', ph['sw'] <= 1 and ph['head'] == 'none' and ph['tall'] < 80, ph)
        pctx.close()
        # P17: Home's Welcome Text shows once per visit, in the middle, and is gone after about 4.5 seconds
        wctx = b.new_context(user_agent=UA, viewport={'width': 1280, 'height': 800}); wctx.add_cookies(cookies); wctx.route('**/*', handle)
        wp = wctx.new_page(); wp.goto(B + '/', wait_until='domcontentloaded'); until(lambda: wp.evaluate("!!document.querySelector('#welcome.in')"), 5000)
        w1 = wp.evaluate("(() => { const e = document.querySelector('#welcome'); return e ? [e.textContent, e.classList.contains('in'), getComputedStyle(document.querySelector('.stage')).opacity] : null })()")
        until(lambda: wp.evaluate("!document.querySelector('#welcome')"), 8000); gone = wp.evaluate("!document.querySelector('#welcome')")   # it fades away by itself after about 4.5 s
        wp.reload(wait_until='domcontentloaded'); wp.wait_for_timeout(600)   # on purpose: P17 checks the welcome does NOT come back, so give it the time it took the first time
        again = wp.evaluate("!!document.querySelector('#welcome.in')")
        check('P17 Home shows "Welcome Home" in the middle with the clock hidden, it fades away, and not again on a reload', w1 and w1[0] == 'Welcome Home' and w1[1] and float(w1[2]) < 0.2 and gone and not again, (w1, gone, again))
        bgnow = wp.evaluate("(() => { const n = document.querySelector('#bg .on'); return n ? (n.currentSrc || n.src) : '' })()")
        check('P18 a new site shows the Starter theme (its own pictures, nothing downloaded from other sites)', '/web/home/starter/' in bgnow and not wp.evaluate("window.PUB.bg.photos.length") and wp.evaluate("window.PUB.bg.theme") == 'starter', bgnow)
        check('P19 Home names its own settings Home Settings (pencil menu and window), and has the theme editor and Import', wp.evaluate("!!document.querySelector('[data-hub=item][data-hub-label=\"Home Settings\"]') && document.querySelector('#settingsSheet h2').textContent === 'Home Settings' && !!document.getElementById('themeSheet') && !!document.getElementById('themeFile')"))
        wctx.close()
        # P20-P28: Home's widgets (1.51). Logged in: the Glass Dock, Quick Add, Done, Settings > Widgets, the Scratch Pad, Recent Tools,
        # the calendar. A visitor gets the clock only, and none of the widgets' data. Home never asks another site for anything.
        hctx = b.new_context(user_agent=UA, viewport={'width': 1280, 'height': 800}); hctx.add_cookies(cookies); hctx.route('**/*', handle)
        hctx.add_init_script("try{sessionStorage.setItem('hl-welcomed','1');localStorage.setItem('tools.v3',JSON.stringify({recent:['pdf-merge','qr-generator','gone-tool']}))}catch(e){}")
        hctx.add_init_script(TRACK_REQUESTS)
        outside = []; herr = []
        hctx.on('request', lambda r: outside.append(r.url) if not r.url.startswith((B, 'data:', 'blob:')) else None)
        hp = hctx.new_page(); hp.on('pageerror', lambda e: herr.append(str(e)[:200]))
        hp.goto(B + '/web/home/', wait_until='domcontentloaded'); until(lambda: hp.evaluate("document.querySelectorAll('#wDock .wdb').length > 0"), 10000); settle(hp)
        dock = hp.evaluate("[...document.querySelectorAll('#wDock .wdb')].map(b => b.dataset.w)")
        check('P20 logged in, Home shows the clock and the Glass Dock with the six widgets that are on at first (Scratch Pad, World Clock and Weather are off)',
              dock == ['today', 'countdown', 'note', 'tools', 'calendar', 'quote'] and hp.evaluate("!document.getElementById('clock').hidden && /\\d/.test(document.getElementById('time').textContent) && !document.getElementById('wDock').hidden && document.getElementById('wCards').hidden"), dock)
        glass = hp.evaluate("[...document.querySelectorAll('.corner button')].filter(b => getComputedStyle(b).display !== 'none').map(b => { const c = getComputedStyle(b); return [b.getAttribute('aria-label') || b.textContent, c.backgroundImage.includes('gradient') && (c.backdropFilter || c.webkitBackdropFilter || 'none') !== 'none'] })")
        check('P20b the corner buttons, the site menu (gear) included, are glass like the widgets (not a plain white button)', len(glass) >= 2 and all(g[1] for g in glass) and any(g[0] == 'Menu' for g in glass), glass)
        hp.click('#wDock [data-w=today]'); hp.wait_for_selector('#wQa', timeout=5000)
        hp.fill('#wQa', 'Pay the gardener today'); hp.press('#wQa', 'Enter')
        rows = lambda: hp.evaluate("[...document.querySelectorAll('#wBody .wrow .wr-n')].map(e => e.textContent)")
        until(lambda: 'Pay the gardener' in rows(), 10000); settle(hp)
        tasks = lambda: hp.evaluate("fetch('../tasks/api.php', {credentials: 'same-origin'}).then(r => r.json()).then(j => j.tasks.map(t => [t.name, t.status]))")
        check('P21 Quick Add in Today adds the task (read like Telegram\'s /add) and it shows in the list', ['Pay the gardener', 'active'] in tasks() and hp.evaluate("document.getElementById('wQa').value") == '', rows())
        hp.click('#wBody .wrow:has-text("Pay the gardener") input'); until(lambda: 'Pay the gardener' not in rows(), 10000); settle(hp)
        check('P22 ticking it in Today marks it done in Tasks', ['Pay the gardener', 'done'] in tasks(), tasks())
        hp.click('#wDock [data-w=calendar]'); hp.wait_for_selector('#wBody .wcal', timeout=5000)
        cal = hp.evaluate("(() => { const d = new Date(), n = new Date(d.getFullYear(), d.getMonth() + 1, 0).getDate(); return [document.querySelectorAll('#wBody .wcal-d').length === n, document.querySelectorAll('#wBody .wcal-d.now').length] })()")
        hp.click('#wBody [data-cal=\"1\"]'); nxt = hp.evaluate("document.querySelectorAll('#wBody .wcal-d.now').length")
        check('P23 Month Calendar shows every day of the month with today marked, and moves to the next month', cal == [True, 1] and nxt == 0, (cal, nxt))
        hp.click('#wDock [data-w=tools]'); hp.wait_for_selector('#wBody .wtools a, #wBody .wmsg', timeout=5000); settle(hp)
        tl = hp.evaluate("[...document.querySelectorAll('#wBody .wtools a')].map(a => [a.textContent, a.getAttribute('href')])")
        check('P24 Recent Tools lists the tools used in this browser by name (a tool that no longer exists is left out), each opening in Tools',
              tl == [['Merge PDF', '../tools/#/tool/pdf-merge'], ['QR Code Generator', '../tools/#/tool/qr-generator']], tl)
        hp.keyboard.press('Escape'); until(lambda: hp.evaluate("document.getElementById('wPanel').hidden"), 3000)
        # Settings > Widgets: Glance and Expand, the Scratch Pad on, a countdown and a pinned note from Info
        hp.evaluate("document.getElementById('editBtn').click()"); hp.wait_for_selector('#ix-ebar:not([hidden])', timeout=10000)
        pill = hp.evaluate("[...document.querySelectorAll('#ix-ebar .ix-eb')].map(b => b.getAttribute('aria-label'))")
        menu = hp.evaluate("[...document.querySelectorAll('[data-hub=item][data-hub-main]')].map(e => e.dataset.hubLabel)")
        check('P36 editing shows the shared edit bar: Background, Widgets, Home Settings and Done (no Log Out); the menu offers only Home Settings for this page',
              pill == ['Background', 'Widgets', 'Home Settings', 'Done'] and menu == ['Home Settings'] and not hp.evaluate("!!document.getElementById('logoutBtn')"), (pill, menu))
        hp.click('#ix-ebar [data-tip="Home Settings"]'); hp.wait_for_selector('#settingsSheet[open]', timeout=5000)
        hs = hp.evaluate("[!!document.querySelector('#settingsSheet #sHomeBar'), document.querySelector('#settingsSheet').textContent.includes('Show the Bar on the Home Page'), !document.querySelector('#settingsSheet #wStyle')]")
        check('P37 Home Settings has Show the Bar on the Home Page, and the widgets are no longer in it', hs == [True, True, True], hs)
        hp.evaluate("document.getElementById('settingsSheet').close()")
        hp.click('#ix-ebar [data-tip="Widgets"]'); hp.wait_for_selector('#widgetsSheet[open]', timeout=5000)
        until(lambda: hp.evaluate("document.querySelectorAll('#wNote optgroup').length > 0"), 10000)
        hp.click('#wStyle [data-v=glance]'); hp.check('#wList [data-wl=pad]')
        hp.fill('#wCdName', 'Lakeside Weekend'); hp.fill('#wCdDate', hp.evaluate("(() => { const tz = (window.HUB && HUB.state && HUB.state().tz) || 'Asia/Kolkata'; const d = new Date(new Intl.DateTimeFormat('en-CA', { timeZone:tz }).format(new Date()) + 'T12:00:00Z'); d.setUTCDate(d.getUTCDate() + 12); return d.toISOString().slice(0, 10); })()"))   # 12 days from today in the site's time zone (not this computer's)
        note = hp.evaluate("(() => { const o = document.querySelector('#wNote optgroup option'); document.getElementById('wNote').value = o.value; return o.textContent })()")
        enc = hp.evaluate("[...document.querySelectorAll('#wNote optgroup')].map(g => g.label)")
        hp.click('#wList [data-wmove=\"5\"][data-by=\"-1\"]')   # Scratch Pad one up
        hp.check('#wList [data-wl=worldclock]'); hp.check('#wList [data-wl=weather]')
        hp.fill('[data-zt=\"0\"]', 'Mars/Nowhere'); hp.click('#saveWidgets'); until(lambda: 'not a time zone' in hp.evaluate("document.getElementById('ix-toast').textContent"), 5000)
        badzone = hp.evaluate("[document.getElementById('widgetsSheet').open, document.getElementById('ix-toast').textContent]")
        hp.fill('[data-zt=\"0\"]', 'Asia/Tokyo'); hp.fill('[data-zn=\"0\"]', 'Aunt Mei'); hp.fill('#wTown', 'Bengaluru')
        # a forecast already saved on the server, so the test asks Open-Meteo nothing (the server side is tested in home_widgets_test.php)
        subprocess.run(['php', '-r', "require '%s/web/home/lib.php'; write_json(weather_file(), ['key' => 'bengaluru|c', 'town' => 'bengaluru', 'place' => ['name' => 'Bengaluru, Karnataka', 'lat' => 12.97, 'lon' => 77.59], 'saved' => time(), 'data' => ['temp' => 27, 'text' => 'Partly cloudy', 'pic' => 'partly', 'day' => true, 'days' => [['date' => date('Y-m-d'), 'max' => 29, 'min' => 19, 'text' => 'Rain', 'pic' => 'rain']], 'place' => 'Bengaluru, Karnataka', 'unit' => '°C', 'saved' => time()]]);" % D], check=True)
        hp.click('#saveWidgets'); until(lambda: not hp.evaluate("document.getElementById('widgetsSheet').open"), 10000)
        hp.click('#ix-ebar [data-ix-done]'); until(lambda: hp.evaluate("document.querySelectorAll('#wCards .wcard').length") == 9, 10000); settle(hp)
        cards = hp.evaluate("[...document.querySelectorAll('#wCards .wcard')].map(c => [c.dataset.w, c.querySelector('.wc-h span').textContent])")
        check('P25 Settings > Widgets switches to Glance and Expand, turns the Scratch Pad on, moves it, and sets the countdown and the note (no encrypted tab offered)',
              [c[0] for c in cards] == ['today', 'countdown', 'note', 'tools', 'pad', 'calendar', 'quote', 'worldclock', 'weather'] and cards[1][1] == 'Lakeside Weekend' and cards[2][1] == note and hp.evaluate("document.getElementById('wDock').hidden")
              and 'Trip Notes' in enc, (cards, note, enc))
        hp.click('#wCards [data-w=countdown]'); hp.wait_for_selector('#wBody .wcount', timeout=5000)
        check('P26 the countdown opens to the days left', hp.evaluate("document.querySelector('#wBody .wcount b').textContent") == '12', hp.evaluate("document.getElementById('wBody').textContent"))
        hp.click('#wCards [data-w=pad]'); hp.wait_for_selector('#wPad', timeout=5000)
        hp.fill('#wPad', 'Buy a spare key'); until(lambda: hp.evaluate("document.getElementById('wPadState').textContent") == 'Saved', 10000)
        hp.reload(wait_until='domcontentloaded'); until(lambda: hp.evaluate("document.querySelectorAll('#wCards .wcard').length") == 6, 10000); settle(hp)
        hp.click('#wCards [data-w=pad]'); hp.wait_for_selector('#wPad', timeout=5000)
        check('P27 the Scratch Pad saves by itself and keeps the text after a reload', hp.evaluate("document.getElementById('wPad').value") == 'Buy a spare key')
        check('P32 World Clock: a name that is not a time zone is refused (Settings stays open), a real one shows with its name and time',
              badzone[0] and 'not a time zone' in badzone[1] and hp.evaluate("(() => { const c = document.querySelector('#wCards [data-w=worldclock]'); return c.querySelector('.wc-h span').textContent === 'Aunt Mei' && /\\d:\\d\\d/.test(c.textContent) })()"), badzone)
        hp.click('#wCards [data-w=weather]'); hp.wait_for_selector('#wBody .wwnow', timeout=5000)
        check('P33 Weather, once turned on with a town, shows the forecast your server saved (27°C, Bengaluru) and credits Open-Meteo',
              hp.evaluate("document.querySelector('#wBody .wwnow b').textContent") == '27°C' and 'Bengaluru, Karnataka' in hp.evaluate("document.getElementById('wBody').textContent") and hp.evaluate("!!document.querySelector('#wBody a[href=\"https://open-meteo.com/\"]')"))
        hp.click('#wCards [data-w=quote]'); hp.wait_for_selector('#wQuoteNext', timeout=5000)
        q1 = hp.evaluate("document.querySelector('#wBody blockquote').textContent"); hp.click('#wQuoteNext'); q2 = hp.evaluate("document.querySelector('#wBody blockquote').textContent")
        check('P34 Quote of the Day shows a quote, and Another Quote shows a different one', len(q1) > 5 and q2 != q1 and hp.evaluate("document.activeElement.id") == 'wQuoteNext', (q1, q2))
        hp.evaluate("document.getElementById('editBtn').click()"); hp.wait_for_selector('#ix-ebar:not([hidden])', timeout=10000)
        hp.click('#ix-ebar [data-tip="Widgets"]'); hp.wait_for_selector('#widgetsSheet[open]', timeout=5000)
        hp.uncheck('#wShow'); hp.click('#saveWidgets'); until(lambda: not hp.evaluate("document.getElementById('widgetsSheet').open"), 10000)
        hp.click('#ix-ebar [data-ix-done]'); settle(hp)
        off = hp.evaluate("[document.querySelectorAll('#wCards .wcard, #wDock .wdb').length, document.body.classList.contains('has-widgets'), !document.getElementById('clock').hidden]")
        check('P38 Widgets > Show Widgets off hides every widget (the clock stays)', off == [0, False, True], off)
        hp.evaluate("document.getElementById('editBtn').click()"); hp.wait_for_selector('#ix-ebar:not([hidden])', timeout=10000)
        hp.click('#ix-ebar [data-tip="Widgets"]'); hp.wait_for_selector('#widgetsSheet[open]', timeout=5000); hp.check('#wShow'); hp.click('#saveWidgets')
        until(lambda: not hp.evaluate("document.getElementById('widgetsSheet').open"), 10000); hp.click('#ix-ebar [data-ix-done]')
        def home_bar(on):
            hp.evaluate("document.getElementById('editBtn').click()"); hp.wait_for_selector('#ix-ebar:not([hidden])', timeout=10000)
            hp.click('#ix-ebar [data-tip="Home Settings"]'); hp.wait_for_selector('#settingsSheet[open]', timeout=5000)
            (hp.check if on else hp.uncheck)('#sHomeBar')
            with hp.expect_navigation(timeout=15000): hp.click('#saveSettings')   # the bar comes or goes with a reload
            until(lambda: hp.evaluate("!!(window.HUB && HUB.state && 'homeBar' in HUB.state())"), 10000); settle(hp)
            return hp.evaluate("[document.documentElement.hasAttribute('data-ix-bar'), HUB.state().homeBar]")
        bar_on = home_bar(True); bar_off = home_bar(False)
        check('P39 Home Settings > Show the Bar on the Home Page turns the bar on and off (the same setting as in Settings)', bar_on == [True, True] and bar_off == [False, False], (bar_on, bar_off))
        hctx.close()
        # P41: an open widget never sits on top of the cards (Glance and Expand, nine cards in two rows) or the dock, and the clock
        # steps aside when there is no room for both (seen on the owner's screen: the panel lay over the cards, unreadable)
        def placed(style, w, h, wid):
            subprocess.run(['php', '-r', "require '%s/web/home/lib.php'; $s = load_site(); $s['widgets']['style'] = '%s'; write_json(SITE_FILE, clean_site($s));" % (D, style)], check=True)
            pc = b.new_context(user_agent=UA, viewport={'width': w, 'height': h}); pc.add_cookies(cookies); pc.route('**/*', handle); pc.add_init_script(TRACK_REQUESTS)
            pc.add_init_script("try{sessionStorage.setItem('hl-welcomed','1')}catch(e){}"); pp = pc.new_page(); pp.goto(B + '/web/home/', wait_until='domcontentloaded')
            sel = '#wCards' if style == 'glance' else '#wDock'
            until(lambda: pp.evaluate("document.querySelectorAll('%s [data-w]').length") == 9, 10000); settle(pp)
            out = []
            for x in wid:
                pp.click('%s [data-w=%s]' % (sel, x)); pp.wait_for_selector('#wPanel:not([hidden])', timeout=5000); settle(pp)
                out.append(pp.evaluate("(() => { const p = document.getElementById('wPanel').getBoundingClientRect(), a = document.querySelector('%s').getBoundingClientRect(), c = document.getElementById('clock'); "
                                       "return (p.bottom <= a.top + 1 ? 1 : 0) + (p.top >= 0 ? 1 : 0) + (getComputedStyle(c).visibility === 'hidden' || c.getBoundingClientRect().bottom <= p.top + 8 ? 1 : 0) })()" % sel))
                pp.click('#wClose')
            pc.close(); return out
        g = placed('glance', 1280, 800, ['quote', 'weather', 'today']); d = placed('dock', 1280, 700, ['weather', 'today'])
        check('P41 an open widget sits above the cards or the dock (never on them), stays on screen, and the clock makes way when there is no room', g == [3, 3, 3] and d == [3, 3], (g, d))
        # P40: the glass. The computer's "reduce transparency" (Windows: Transparency effects off) no longer turns it into a dark block;
        # Home Settings > Glass > Tint decides: Light (see-through, the default), Dark, or Solid.
        def glass(tint=None, reduce=False):
            if tint: subprocess.run(['php', '-r', "require '%s/web/home/lib.php'; $s = load_site(); $s['panel']['tint'] = '%s'; write_json(SITE_FILE, clean_site($s));" % (D, tint)], check=True)
            gc = b.new_context(user_agent=UA, viewport={'width': 1280, 'height': 800}); gc.add_cookies(cookies); gc.route('**/*', handle); gc.add_init_script(TRACK_REQUESTS)
            gc.add_init_script("try{sessionStorage.setItem('hl-welcomed','1')}catch(e){}"); gp = gc.new_page()
            if reduce: gc.new_cdp_session(gp).send('Emulation.setEmulatedMedia', {'features': [{'name': 'prefers-reduced-transparency', 'value': 'reduce'}]}) if BROWSER == 'chromium' else None
            gp.goto(B + '/web/home/', wait_until='domcontentloaded'); until(lambda: gp.evaluate("document.querySelectorAll('.wcard, .wdb').length > 0"), 10000); settle(gp)
            el = '.wcard' if gp.evaluate("!!document.querySelector('.wcard')") else '#wDock'
            r = gp.evaluate("(() => { const c = getComputedStyle(document.querySelector('%s')); const m = c.backgroundColor.match(/rgba?\\(([^)]+)\\)/)[1].split(',').map(Number); return [m[0], m[3] === undefined ? 1 : m[3], (c.backdropFilter || c.webkitBackdropFilter || 'none') !== 'none'] })()" % el)
            gc.close(); return r
        light_reduce = glass('light', reduce=True); dark = glass('dark'); solid = glass('solid'); glass('light')
        check('P40 Glass Tint: Light stays see-through and blurred even when the computer asks to reduce transparency; Dark is a dark see-through tint; Solid is not see-through',
              light_reduce[0] == 255 and light_reduce[1] < 0.5 and light_reduce[2] and dark[0] < 60 and dark[1] < 0.5 and dark[2] and solid[1] > 0.85 and not solid[2], (light_reduce, dark, solid))
        # P48: Settings > Overview after an update shows What's New from a small answer (the newest version only, 1.53.1)
        wctx = b.new_context(user_agent=UA, viewport={'width': 1280, 'height': 860}); wctx.add_cookies(cookies); wctx.route('**/*', handle); wctx.add_init_script(TRACK_REQUESTS)
        wsz = []; wp = wctx.new_page(); wp.on('requestfinished', lambda r: wsz.append((r.url, len(r.response().body() or b''))) if 'action=changelog' in r.url else None)
        wp.goto(B + '/web/settings/', wait_until='domcontentloaded'); until(lambda: wp.evaluate("!!document.querySelector('.ix-new')"), 15000, 200)
        wn = wp.evaluate("(document.querySelector('.ix-new .ix-ec-t') || {}).textContent || ''")
        wctx.close()
        check("P48 Settings > Overview shows What's New for this version, loading only its notes (not the whole changelog)",
              ver in wn and wsz and all('latest=1' in u and n < 20000 for u, n in wsz), (wn, [(u[-40:], n) for u, n in wsz]))
        # P49: the shared encryption of encrypted tabs (shell.js HubVault, 1.53.2; before, Links, Info and Settings each had a copy).
        # It opens what the old copies made (tests/fixtures/vault-vectors.json), and what it makes opens with plain WebCrypto
        # written here, apart from the site's code: so the stored format did not change.
        vec = json.load(open(os.path.join(ROOT, 'tests', 'fixtures', 'vault-vectors.json'), encoding='utf-8'))
        xctx = b.new_context(user_agent=UA); xctx.route('**/*', handle); xp = xctx.new_page()
        xp.goto(B + '/web/links/', wait_until='domcontentloaded'); until(lambda: xp.evaluate("!!window.HubVault"), 10000)
        vr = xp.evaluate("""async (x) => {
          const V = HubVault, te = new TextEncoder(), td = new TextDecoder(), same = (a, b) => !!a && a.length === b.length && a.every((n, i) => n === b[i]);
          const ub = s => Uint8Array.from(atob(s), c => c.charCodeAt(0));
          const kek = async (pass, salt, iter) => crypto.subtle.deriveKey({ name: 'PBKDF2', salt, iterations: iter, hash: 'SHA-256' },
            await crypto.subtle.importKey('raw', te.encode(pass), 'PBKDF2', false, ['deriveKey']), { name: 'AES-GCM', length: 256 }, false, ['decrypt']);
          const gcm = async (key, b) => new Uint8Array(await crypto.subtle.decrypt({ name: 'AES-GCM', iv: b.subarray(0, 12) }, key, b.subarray(12)));
          const r = {}, raw = Uint8Array.from(x.raw);
          r.pass = same(await V.unwrap(x.vault, x.pass), raw);
          r.rec = same(await V.unwrap(x.vault, ' ' + x.rec.toUpperCase().replace(/-/g, ' ') + ' '), raw);   // typed with spaces and capitals
          r.wrong = (await V.unwrap(x.vault, 'not it')) === null && (await V.unwrap({ salt: 'AA==', iter: 1, wrap: '!!' }, 'x')) === null;
          const k = await V.useKey(raw), kd = await V.useKey(raw, true);
          r.open = (await V.unseal(k, x.sealed)) === x.plain && (await V.openDeep(kd, x.sealed)) === x.plain;
          r.deep = (await V.openDeep(k, await V.seal(k, await V.seal(k, x.plain)))) === x.plain && (await V.openDeep(k, 'plain')) === 'plain';
          r.bad = await V.unseal(await V.useKey(new Uint8Array(32)), x.sealed).then(() => false, () => true);
          r.readOnly = await V.seal(kd, 'x').then(() => false, () => true);
          const w = await V.wrap(raw, 'New pass 9', true);
          r.wrapShape = w.v === 1 && w.iter === 600000 && w.login === true && ub(w.salt).length === 16 && !('login' in await V.wrap(raw, 'p', false));
          r.wrapIndep = same(await gcm(await kek('New pass 9', ub(w.salt), w.iter), ub(w.wrap)), raw);
          const s = await V.seal(k, x.plain), raw2 = await crypto.subtle.importKey('raw', raw, 'AES-GCM', false, ['decrypt']);
          r.sealIndep = s.startsWith('enc1:') && s !== await V.seal(k, x.plain) && td.decode(await gcm(raw2, ub(s.slice(5)))) === x.plain;
          await V.devSave('p49', raw); const back = await V.devLoad('p49'); await V.devForget('p49');
          r.device = same(back, raw) && (await V.devLoad('p49')) === null;
          return r;
        }""", vec)
        xctx.close()
        check('P49 the shared encryption opens tabs made before 1.53.2 (passphrase, recovery key typed loosely, text encrypted twice), refuses a wrong key, and what it makes opens with plain WebCrypto',
              vr and all(vr.values()), vr)
        # P50-P54: the edit bar (1.53.5). The pencil menu has only Edit This Page and the page's Settings; while editing, the page's
        # tools sit in one shared pill at the bottom (Home's look), each page's Settings picks Icons Only or Icons and Text, and
        # Settings has no pencil at all.
        ectx = b.new_context(user_agent=UA, viewport={'width': 1280, 'height': 800}); ectx.add_cookies(cookies); ectx.route('**/*', handle); ectx.add_init_script(TRACK_REQUESTS)
        ep = ectx.new_page(); eerr = []; ep.on('pageerror', lambda e: eerr.append(str(e)[:200]))
        def emenu(app):
            ep.goto(B + '/web/%s/' % app, wait_until='domcontentloaded'); until(lambda: ep.evaluate("!!(window.HUB && HUB.state().admin) && document.readyState === 'complete'"), 15000); settle(ep)
            ep.click('#ix-editb'); until(lambda: ep.evaluate("!!document.querySelector('#ix-pop')"), 3000)
            return ep.evaluate("[...document.querySelectorAll('#ix-pop .ix-pl')].map(x => x.textContent)")
        def ebar(): return ep.evaluate("(() => { const e = document.getElementById('ix-ebar'); return e && !e.hidden ? [...e.querySelectorAll('.ix-eb')].map(b => b.getAttribute('aria-label')) : null })()")
        want = {'links': ['Colors', 'Fetch Icons', 'Import Bookmarks', 'Links Settings', 'Done'], 'info': ['Export', 'Info Settings', 'Done'],
                'cards': ['Add Card', 'Holders', 'Banks', 'Export', 'Cards Settings', 'Done'], 'tasks': ['Add Task', 'Export', 'Tasks Settings', 'Done']}
        menus, bars, old, done = {}, {}, {}, {}
        for app in want:
            menus[app] = emenu(app)
            ep.click('#ix-pop button'); until(lambda: ebar() is not None, 5000)
            bars[app] = ebar()
            old[app] = ep.evaluate("[...document.querySelectorAll('[data-hub-ebar]')].every(e => getComputedStyle(e).display === 'none')")
            ep.click('#ix-ebar [data-ix-done]'); until(lambda: ebar() is None, 5000)
            done[app] = not ep.evaluate("document.body.classList.contains('editing')")
        check('P50 the pencil menu offers only Edit This Page and the page\'s own Settings (Links, Info, Cards, Tasks)',
              all(menus[a] == ['Edit This Page', a.title() + ' Settings'] for a in want), menus)
        check('P51 Edit This Page shows the shared edit bar with the page\'s tools, its Settings and Done, the old bar is hidden, and Done ends editing (also on Cards and Tasks, which have no edit mode of their own)',
              bars == want and all(old.values()) and all(done.values()), (bars, old, done))
        # the bar's buttons do what the old ones did
        emenu('links'); ep.click('#ix-pop button'); until(lambda: ebar() is not None, 5000)
        ep.hover('#ix-ebar [data-tip="Colors"]'); until(lambda: ep.evaluate("!!document.getElementById('ix-ebtip')"), 3000)
        tipname = ep.evaluate("document.getElementById('ix-ebtip').textContent")
        ep.click('#ix-ebar [data-tip="Colors"]'); until(lambda: ep.evaluate("document.getElementById('colorDlg').open"), 5000)
        ep.evaluate("document.getElementById('colorDlg').close()")
        ep.click('#ix-ebar [data-tip="Links Settings"]'); until(lambda: ep.evaluate("document.getElementById('setDlg').open"), 5000)
        before = ep.evaluate("document.getElementById('sEditBar').value")
        ep.select_option('#sEditBar', 'text'); ep.click('#setDlg button[value=cancel]'); ep.wait_for_timeout(600)
        cancelled = ep.evaluate("[document.getElementById('ix-ebar').className, (HUB.state().editBar || {}).links || '']")
        ep.click('#ix-ebar [data-tip="Links Settings"]'); until(lambda: ep.evaluate("document.getElementById('setDlg').open"), 5000)
        reopened = ep.evaluate("document.getElementById('sEditBar').value")
        ep.select_option('#sEditBar', 'text'); ep.click('#setSave'); until(lambda: ep.evaluate("document.getElementById('ix-ebar').classList.contains('ix-ebar-text')"), 8000)
        labels = ep.evaluate("[...document.querySelectorAll('#ix-ebar .ix-ebl')].filter(l => l.offsetWidth > 0).map(l => l.textContent)")
        check('P52 Icons Only (the default) shows a button\'s name when you point at it; Links Settings > Edit Bar > Icons and Text writes the names on the bar once saved (Cancel keeps it as it was)',
              tipname == 'Colors' and before == 'icons' and cancelled == ['', ''] and reopened == 'icons' and labels == ['Colors', 'Fetch Icons', 'Import Bookmarks', 'Links Settings'], (tipname, before, cancelled, reopened, labels))
        ep.click('#ix-ebar [data-ix-done]')
        emenu('links'); ep.click('#ix-pop button'); until(lambda: ebar() is not None, 5000)
        kept = ep.evaluate("document.getElementById('ix-ebar').classList.contains('ix-ebar-text')"); ep.click('#ix-ebar [data-ix-done]')
        emenu('tasks'); ep.click('#ix-pop button'); until(lambda: ebar() is not None, 5000)
        other = ep.evaluate("document.getElementById('ix-ebar').className"); ep.click('#ix-ebar [data-ix-done]')
        emenu('cards'); ep.click('#ix-pop button'); until(lambda: ebar() is not None, 5000)
        ep.click('#ix-ebar [data-tip="Export"]'); until(lambda: ep.evaluate("!!document.querySelector('#ix-pop')"), 3000)
        exp = ep.evaluate("[...document.querySelectorAll('#ix-pop .ix-pl')].map(x => x.textContent)"); ep.keyboard.press('Escape'); ep.click('#ix-ebar [data-ix-done]')
        ep.goto(B + '/web/settings/', wait_until='domcontentloaded'); until(lambda: ep.evaluate("!!(window.HUB && HUB.state().admin)"), 15000); settle(ep)
        nopen = ep.evaluate("!document.getElementById('ix-editb') && !document.querySelector('[data-ix=edit]')")
        check('P53 the choice is kept on the server (Links still shows names after a reload) and is per page (Tasks keeps icons); Cards\' Export opens its two exports; Settings has no pencil',
              kept and other == '' and exp == ['PDF (print)', 'Excel Spreadsheet'] and nopen, (kept, other, exp, nopen))
        ectx.close()
        mctx2 = b.new_context(user_agent=UA, viewport={'width': 390, 'height': 780}); mctx2.add_cookies(cookies); mctx2.route('**/*', handle); mctx2.add_init_script(TRACK_REQUESTS)
        mp2 = mctx2.new_page(); mp2.goto(B + '/web/links/', wait_until='domcontentloaded'); until(lambda: mp2.evaluate("!!(window.HUB && HUB.state().admin) && document.readyState === 'complete'"), 15000); settle(mp2)
        mp2.evaluate("document.querySelector('[data-hub=toggle]').click()"); until(lambda: mp2.evaluate("!!document.getElementById('ix-ebar') && !document.getElementById('ix-ebar').hidden"), 5000); settle(mp2)
        ph = mp2.evaluate("(() => { const e = document.getElementById('ix-ebar'), r = e.getBoundingClientRect(); return [r.left >= 0 && r.right <= innerWidth, [...e.querySelectorAll('.ix-ebl')].every(l => l.offsetWidth === 0), e.querySelectorAll('.ix-eb').length] })()")
        mctx2.close()
        check('P54 on a phone the edit bar fits the screen and shows icons only (even with Icons and Text chosen), with no page error', ph == [True, True, 5] and not eerr, (ph, eerr))
        # P55-P58 (1.53.6): no login box for someone logged in who presses the pencil at once; the two-item menu has no line;
        # Cards' one Holders window; the tidied Settings windows of Links, Info and Tasks
        rctx2 = b.new_context(user_agent=UA, viewport={'width': 1280, 'height': 800}); rctx2.add_cookies(cookies); rctx2.route('**/*', handle)
        rctx2.add_init_script("(() => { const f = window.fetch; window.fetch = function (u, o) { if (String(u).includes('action=state')) return new Promise(r => setTimeout(r, 2500)).then(() => f(u, o)); return f(u, o) } })()")
        rp2 = rctx2.new_page(); rp2.goto(B + '/web/info/', wait_until='domcontentloaded'); rp2.wait_for_selector('#ix-editb', timeout=10000)
        early = rp2.evaluate("HUB.state().admin === undefined"); rp2.click('#ix-editb')
        until(lambda: rp2.evaluate("!!document.querySelector('#ix-pop') || !!document.querySelector('#ix-dlg.open')"), 8000)
        race = rp2.evaluate("[!!document.querySelector('#ix-dlg.open'), [...document.querySelectorAll('#ix-pop .ix-pl')].map(x => x.textContent)]"); rctx2.close()
        check('P55 logged in, pressing the pencil before the page has heard from the server waits for it: the menu opens, no login box', early and race == [False, ['Edit This Page', 'Info Settings']], (early, race))
        hctx = b.new_context(user_agent=UA, viewport={'width': 1280, 'height': 860}); hctx.add_cookies(cookies); hctx.route('**/*', handle); hp2 = hctx.new_page(); herr2 = []; hp2.on('pageerror', lambda e: herr2.append(str(e)[:200]))
        hp2.goto(B + '/web/cards/', wait_until='domcontentloaded'); until(lambda: hp2.evaluate("!!(window.HUB && HUB.state().admin) && document.readyState === 'complete'"), 15000); settle(hp2)
        hp2.click('#ix-editb'); until(lambda: hp2.evaluate("!!document.querySelector('#ix-pop')"), 3000)
        lines = hp2.evaluate("document.querySelectorAll('#ix-pop hr').length")
        check('P56 the pencil menu (Edit This Page, Cards Settings) has no dividing line', lines == 0, lines)
        hp2.click('#ix-pop button'); until(lambda: hp2.evaluate("!!document.getElementById('ix-ebar') && !document.getElementById('ix-ebar').hidden"), 5000)
        hrows = lambda: hp2.evaluate("[...document.querySelectorAll('#hlist .hrow')].map(r => [r.querySelector('input').value, r.querySelector('.hn').textContent, r.querySelector('.hdel').disabled])")
        hp2.click('#ix-ebar [data-tip="Holders"]'); until(lambda: hp2.evaluate("document.getElementById('hdlg').open"), 5000)
        first = hrows()
        hp2.fill('#hname', 'Riya'); hp2.click('#hadd'); hp2.click('#hcancel'); hp2.wait_for_timeout(300)
        hp2.click('#ix-ebar [data-tip="Holders"]'); until(lambda: hp2.evaluate("document.getElementById('hdlg').open"), 5000)
        cancelled = [r[0] for r in hrows()]
        who = first[0][0]
        hp2.fill('#hlist input[data-hi="0"]', who + ' K'); hp2.fill('#hname', 'Riya'); hp2.keyboard.press('Enter'); hp2.fill('#hname', 'riya'); hp2.click('#hadd')
        dupmsg = hp2.evaluate("document.getElementById('hmsg').textContent")
        hp2.click('#hok'); until(lambda: not hp2.evaluate("document.getElementById('hdlg').open"), 5000); settle(hp2)
        hp2.click('#ix-ebar [data-tip="Holders"]'); until(lambda: hp2.evaluate("document.getElementById('hdlg').open"), 5000)
        after = hrows()
        hp2.click('#hlist .hrow:last-child .hdel'); hp2.fill('#hlist input[data-hi="0"]', who); hp2.click('#hok'); until(lambda: not hp2.evaluate("document.getElementById('hdlg').open"), 5000); settle(hp2)
        hp2.click('#ix-ebar [data-tip="Holders"]'); until(lambda: hp2.evaluate("document.getElementById('hdlg').open"), 5000)
        back = [r[0] for r in hrows()]; hp2.click('#hcancel')
        check('P57 Cards: one Holders window lists every holder with their cards (a holder with cards cannot be deleted); renaming moves the cards with it, Add Holder adds one (not twice), Delete removes an empty one, Cancel keeps everything as it was',
              all(r[1].endswith('cards') or r[1] == '1 card' for r in first) and any(r[2] for r in first) and cancelled == [r[0] for r in first]
              and after[0][0] == who + ' K' and after[0][1] == first[0][1] and after[-1] == ['Riya', '0 cards', False] and 'already' in dupmsg
              and back == [r[0] for r in first] and not herr2, (first, cancelled, after, dupmsg, back, herr2))
        def sdlg(app, dlg):
            hp2.goto(B + '/web/%s/' % app, wait_until='domcontentloaded'); until(lambda: hp2.evaluate("!!(window.HUB && HUB.state().admin) && document.readyState === 'complete'"), 15000); settle(hp2)
            hp2.click('#ix-editb'); until(lambda: hp2.evaluate("!!document.querySelector('#ix-pop')"), 3000); hp2.click('#ix-pop button:nth-child(2)')
            until(lambda: hp2.evaluate("document.getElementById('%s').open" % dlg), 8000); settle(hp2)
            return hp2.evaluate("""(() => { const d = document.getElementById('%s');
              return [[...d.querySelectorAll('h3, h4')].map(h => h.textContent.trim()), [...d.querySelectorAll('.ix-srow')].map(r => [r.firstElementChild.firstChild.textContent.trim(), r.lastElementChild.tagName, r.lastElementChild.textContent.trim().slice(0, 12), r.getBoundingClientRect().width > 0 ? Math.round(r.lastElementChild.getBoundingClientRect().width) : 0])] })()""" % dlg)
        ls_, is_ = sdlg('links', 'setDlg'), sdlg('info', 'setdlg')
        ts_ = sdlg('tasks', 'sdlg'); tl = hp2.evaluate("[...document.querySelectorAll('#sdlg .sh3')].map(h => h.textContent.trim()); ")
        lookhas = hp2.evaluate("(() => { const h = [...document.querySelectorAll('#sdlg .sh3')].pop(); let n = h.nextElementSibling; return !!n.querySelector('#s-editbar') && !!n.querySelector('#s-week') })()")
        hctx.close()
        rows_ok = lambda r: [x[:2] for x in r[1]] == [['Text Sizes', 'BUTTON'], ['Edit Bar', 'SELECT']] and r[1][0][2] == 'Change' and r[1][0][3] < 200
        check('P58 Links and Info Settings: a Look section with Text Sizes (a small Change button) and Edit Bar side by side with their help; Tasks Settings: General, When the Page Opens and Look, with Edit Bar in Look',
              rows_ok(ls_) and rows_ok(is_) and 'Look' in ls_[0] and 'Look' in is_[0] and tl == ['General', 'When the Page Opens', 'Look'] and lookhas, (ls_, is_, tl, lookhas))
        # P60-P62 (1.53.8): the login window; Cards Settings like Links and Info; the Banks window and its own icon
        lctx2 = b.new_context(user_agent=UA, viewport={'width': 1280, 'height': 800}); lctx2.route('**/*', handle)
        lp2 = lctx2.new_page(); lp2.goto(B + '/web/links/', wait_until='domcontentloaded'); until(lambda: lp2.evaluate("!!(window.HUB && HUB.state().admin === false)"), 15000)
        lp2.evaluate("setTimeout(() => HUB.login(), 0); 0"); lp2.wait_for_selector('#ix-dlg.open #ix-pw', timeout=5000)
        lw = lp2.evaluate("""(() => { const d = document.getElementById('ix-dlg'); return [!!d.querySelector('.ix-llogo'), d.querySelector('h2').textContent, !/Welcome!/.test(d.textContent),
          !document.getElementById('ix-ltrust').checked, document.getElementById('ix-lg').offsetWidth >= document.getElementById('ix-pw').offsetWidth - 2] })()""")
        lp2.fill('#ix-pw', 'not the password'); lp2.click('#ix-eye'); shownpw = lp2.evaluate("[document.getElementById('ix-pw').type, document.getElementById('ix-eye').getAttribute('aria-pressed')]")
        lp2.click('#ix-eye'); hiddenpw = lp2.evaluate("document.getElementById('ix-pw').type")
        lp2.click('#ix-lg'); until(lambda: lp2.evaluate("document.getElementById('ix-le').textContent.length > 0"), 8000)
        lerr = lp2.evaluate("(() => { const e = document.getElementById('ix-le'), f = document.getElementById('ix-pw'); return [e.textContent, e.getBoundingClientRect().top > f.getBoundingClientRect().bottom, e.getBoundingClientRect().top < document.getElementById('ix-lg').getBoundingClientRect().top] })()")
        lctx2.close()
        check('P60 the login window shows the site logo and "Log In to <site name>", a password you can show and hide, a wide Log In button, a wrong password\'s message right under the field, and Trust This Browser not ticked',
              lw[0] and lw[1].startswith('Log In to ') and lw[2] and lw[3] and lw[4] and shownpw == ['text', 'true'] and hiddenpw == 'password' and lerr[0] and lerr[1] and lerr[2], (lw, shownpw, hiddenpw, lerr))
        cctx = b.new_context(user_agent=UA, viewport={'width': 1280, 'height': 860}); cctx.add_cookies(cookies); cctx.route('**/*', handle); cp = cctx.new_page(); cerr = []; cp.on('pageerror', lambda e: cerr.append(str(e)[:200]))
        cp.goto(B + '/web/cards/', wait_until='domcontentloaded'); cp.wait_for_selector('#list .row.item', timeout=15000); settle(cp)
        cp.click('#ix-editb'); until(lambda: cp.evaluate("!!document.querySelector('#ix-pop')"), 3000); cp.click('#ix-pop button'); until(lambda: cp.evaluate("!!document.getElementById('ix-ebar') && !document.getElementById('ix-ebar').hidden"), 5000)
        icons = cp.evaluate("[...document.querySelectorAll('#ix-ebar .ix-eb')].map(b => [b.getAttribute('aria-label'), b.querySelector('path').getAttribute('d')])")
        idict = dict(icons)
        cp.click('#ix-ebar [data-tip="Cards Settings"]'); until(lambda: cp.evaluate("document.getElementById('tdlg').open"), 5000)
        cs = cp.evaluate("[[...document.querySelectorAll('#tdlg h3')].map(h => h.textContent), [...document.querySelectorAll('#tdlg .ix-srow')].map(r => [r.firstElementChild.firstChild.textContent.trim(), r.lastElementChild.tagName])]")
        cp.evaluate("document.getElementById('tdlg').close()")
        banks_of = lambda: cp.evaluate("[...new Set([...document.querySelectorAll('#list .row.item')].map(r => r.textContent))].join('|')")
        cp.click('#ix-ebar [data-tip="Banks"]'); until(lambda: cp.evaluate("document.getElementById('bdlg').open"), 5000)
        brows = cp.evaluate("[...document.querySelectorAll('#blist .brow')].map(r => [r.querySelector('[data-bf=name]').value, r.querySelector('.hn').textContent])")
        cp.fill('#blist .brow:first-child [data-bf=name]', 'Lakeside Bank'); cp.click('#bcancel'); cp.wait_for_timeout(300)
        kept = 'Lakeside Bank' not in banks_of()
        cp.click('#ix-ebar [data-tip="Banks"]'); until(lambda: cp.evaluate("document.getElementById('bdlg').open"), 5000)
        old = brows[0][0]
        cp.fill('#blist .brow:first-child [data-bf=name]', 'Lakeside Bank'); cp.fill('#blist .brow:first-child [data-bf=url]', 'not a link'); cp.click('#bsave'); cp.wait_for_timeout(300)
        badlink = cp.evaluate("[document.getElementById('bdlg').open, document.getElementById('bmsg').textContent]")
        cp.fill('#blist .brow:first-child [data-bf=url]', 'https://lakeside.example/'); cp.click('#bsave'); until(lambda: not cp.evaluate("document.getElementById('bdlg').open"), 5000); settle(cp)
        renamed = banks_of(); link = cp.evaluate("[...document.querySelectorAll('#list a[href]')].some(a => a.href === 'https://lakeside.example/')")
        cp.click('#ix-ebar [data-tip="Banks"]'); until(lambda: cp.evaluate("document.getElementById('bdlg').open"), 5000)
        cp.fill('#blist .brow:first-child [data-bf=name]', old); cp.fill('#blist .brow:first-child [data-bf=url]', ''); cp.click('#bsave'); until(lambda: not cp.evaluate("document.getElementById('bdlg').open"), 5000)
        cp.click('#ix-ebar [data-tip="Holders"]'); until(lambda: cp.evaluate("document.getElementById('hdlg').open"), 5000); cp.focus('#hlist input')
        ring = cp.evaluate("(() => { const s = getComputedStyle(document.querySelector('#hlist input')); return [s.outlineStyle, s.borderColor] })()"); cp.click('#hcancel')
        cctx.close()
        check('P61 Cards Settings looks like Links and Info: Page, then Look with Text Sizes, Edit Bar and Column Order rows',
              cs == [['Page', 'Look'], [['Text Sizes', 'BUTTON'], ['Edit Bar', 'SELECT'], ['Column Order', 'BUTTON']]], cs)
        check('P62 Cards: Banks has its own icon (not Add Card\'s) and one window like Holders: every bank with its cards, a rename moves the cards, a link is kept, a wrong link is refused, Cancel changes nothing; a holder\'s name box gets a clean border, not a cut-off outline',
              idict.get('Banks') and idict.get('Banks') != idict.get('Add Card') and brows and all(x[1].endswith('card') or x[1].endswith('cards') for x in brows) and kept
              and badlink[0] and 'https://' in badlink[1] and 'Lakeside Bank' in renamed and old not in renamed and link and ring[0] == 'none' and not cerr, (icons, brows, kept, badlink, link, ring, cerr))
        # P47: Home's refresh button (1.53). The next background is ready, so a press swaps at once; three quick presses on a slow
        # network give one change (it never starts over); when nothing arrives it gives up within about 20 s and the button works again.
        rctx = b.new_context(user_agent=UA, viewport={'width': 1280, 'height': 800}); rctx.route('**/*', handle); rctx.add_init_script(TRACK_REQUESTS)
        rctx.add_init_script("try{sessionStorage.setItem('hl-welcomed','1')}catch(e){}")
        net = {'delay': 0, 'hang': False}
        def bg_route(r):
            if r.request.resource_type in ('image', 'media') and '/hub/media/' not in r.request.url and ('starter' in r.request.url or '/photos/' in r.request.url or '/uploads/' in r.request.url):
                if net['hang']: return   # never answered
                if net['delay']: time.sleep(net['delay'])
            r.fallback()
        rp = rctx.new_page(); rp.route('**/*', bg_route)
        rp.goto(B + '/web/home/', wait_until='domcontentloaded')
        cur = lambda: rp.evaluate("(() => { const n = document.querySelector('#bg > .on'); return n ? (n.currentSrc || n.src || n.poster || '') : '' })()")
        until(cur, 15000); rp.wait_for_timeout(2500)   # the next one is fetched about 0.6 s after the first appears
        first = cur(); t0 = time.time(); rp.click('#nextBg'); until(lambda: cur() != first, 5000); quick = time.time() - t0
        rp.wait_for_timeout(300)
        net['delay'] = 2.5; second = cur(); t0 = time.time()
        rp.evaluate("document.getElementById('nextBg').click()")
        changes = []; rp.wait_for_timeout(150)
        dimmed = rp.evaluate("document.getElementById('bg').classList.contains('wait') || document.querySelector('#bg > .on') !== null && getComputedStyle(document.querySelector('#bg > .on')).filter !== 'none'") or cur() != second   # (or it swapped at once: that one was ready)
        rp.evaluate("document.getElementById('nextBg').click(); document.getElementById('nextBg').click()")
        end = time.time() + 9
        while time.time() < end:
            c = cur()
            if c != second and (not changes or changes[-1] != c): changes.append(c)
            if changes and time.time() - t0 > 6: break
            rp.wait_for_timeout(100)
        slow = time.time() - t0
        net['delay'] = 0; net['hang'] = True
        rp.wait_for_timeout(1500); rp.evaluate("document.getElementById('nextBg').click()")   # uses up the one that was ready...
        until(lambda: not rp.evaluate("document.getElementById('nextBg').classList.contains('spin')"), 20000, 250); rp.wait_for_timeout(1500)   # ...and the next one's download hangs
        rp.evaluate("document.getElementById('nextBg').click()"); t0 = time.time()
        freed = until(lambda: not rp.evaluate("document.getElementById('bg').classList.contains('wait') || document.getElementById('nextBg').classList.contains('spin')"), 25000, 250)
        gave_up = time.time() - t0
        net['hang'] = False; rctx.close()
        check('P47 Home refresh: a ready background swaps at once; on a slow network the picture dims (or swaps) at once and three quick presses never stack up; when nothing arrives the button frees itself within about 20 s',
              first and quick < 1.2 and 1 <= len(changes) <= 2 and dimmed and freed and gave_up < 22 and cur is not None, (first[-30:], round(quick, 2), [c[-25:] for c in changes], round(slow, 1), dimmed, freed, round(gave_up, 1)))
        # P35: a phone, logged in, all nine widgets in the Glass Dock: every dock button fits on the screen, and an open widget
        # does not sit on top of the chips (on a short screen they make way)
        subprocess.run(['php', '-r', "require '%s/web/home/lib.php'; $s = load_site(); $s['widgets']['style'] = 'dock'; write_json(SITE_FILE, clean_site($s));" % D], check=True)
        mctx = b.new_context(user_agent=UA, viewport={'width': 390, 'height': 700}); mctx.add_cookies(cookies); mctx.route('**/*', handle); mctx.add_init_script(TRACK_REQUESTS)
        mctx.add_init_script("try{sessionStorage.setItem('hl-welcomed','1')}catch(e){}")
        mp = mctx.new_page(); mp.goto(B + '/web/home/', wait_until='domcontentloaded'); until(lambda: mp.evaluate("document.querySelectorAll('#wDock .wdb').length") == 9, 10000); settle(mp)
        fit = mp.evaluate("(() => { const d = document.getElementById('wDock'); return [d.scrollWidth <= d.clientWidth + 1, [...d.children].every(b => { const r = b.getBoundingClientRect(); return r.left >= 0 && r.right <= innerWidth }) ] })()")
        mp.click('#wDock [data-w=today]'); mp.wait_for_selector('#wQa', timeout=5000)
        chips = mp.evaluate("getComputedStyle(document.getElementById('wChips')).visibility")
        check('P35 on a phone all nine dock buttons fit on the screen, and an open widget does not cover the chips (they make way)', fit == [True, True] and chips == 'hidden', (fit, chips))
        mctx.close()
        # a visitor: the clock, no widgets, and nothing private in the page or asked of the server
        vctx = b.new_context(user_agent=UA, viewport={'width': 390, 'height': 844}); vctx.route('**/*', handle); vctx.add_init_script(TRACK_REQUESTS)
        vctx.add_init_script("try{sessionStorage.setItem('hl-welcomed','1')}catch(e){}")
        vreq = []; vctx.on('request', lambda r: vreq.append(r.url.replace(B, '')))
        vctx.on('request', lambda r: outside.append(r.url) if not r.url.startswith((B, 'data:', 'blob:')) else None)
        vp = vctx.new_page(); vp.on('pageerror', lambda e: herr.append(str(e)[:200])); vp.goto(B + '/web/home/', wait_until='domcontentloaded'); settle(vp)
        html = vp.content()
        v = vp.evaluate("[window.PUB.widgets, !document.getElementById('clock').hidden, document.body.classList.contains('has-widgets'), document.documentElement.scrollWidth - innerWidth]")
        asked = [u for u in vreq if any(k in u for k in ('/tasks/api.php', '/info/api.php', 'action=pad', 'tools-list.json'))]
        check('P28 a visitor (on a phone) sees the clock and no widgets: no widget settings in the page, no Tasks, Info or Scratch Pad request, nothing scrolls sideways',
              v[0] is None and v[1] and not v[2] and v[3] <= 1 and 'Lakeside Weekend' not in html and not asked, (v, asked))
        check('P29 Home asks no other site for anything (logged in or not), with no page error', not outside and not herr, (outside[:3], herr[:2]))
        vctx.close()
        # P30-P31: Links Settings > Icons from Google (1.51). Off at first: a link with no icon saved on the server shows its letter,
        # and the browser asks Google nothing (also not for an old saved service address). On: the browser asks the service.
        def links_data(web):
            code = ("require '%s/web/hub/core/lib/boot.php'; hub_write_guarded(hub_dir('links') . '/bookmarks.php', ['settings' => %s, 'tabs' => [['id' => 'daily', 'title' => 'Daily', "
                    "'categories' => [['id' => 'c1', 'title' => 'News', 'col' => 0, 'links' => [['title' => 'Example', 'url' => 'https://example.org/'], "
                    "['title' => 'Old Icon', 'url' => 'https://example.net/', 'icon' => 'https://www.google.com/s2/favicons?domain=example.net&sz=64']]]]]]]);") % (D, "['webIcons' => true]" if web else '[]')
            subprocess.run(['php', '-r', code], check=True)
        lctx = b.new_context(user_agent=UA, viewport={'width': 1280, 'height': 800}); lctx.route('**/*', handle); lctx.add_init_script(TRACK_REQUESTS)
        lout = []; lctx.on('request', lambda r: lout.append(r.url) if not r.url.startswith((B, 'data:', 'blob:')) else None)
        lp = lctx.new_page(); links_data(False); lp.goto(B + '/web/links/', wait_until='domcontentloaded'); lp.wait_for_selector('.ico', timeout=10000); settle(lp)
        off = (lp.evaluate("[...document.querySelectorAll('.ico')].map(e => e.textContent + (e.querySelector('img') ? '+img' : ''))"), list(lout))
        check('P30 Links: Icons from Google is off at first, so links without a saved icon show their letter and the browser asks Google nothing', off[0] == ['E', 'O'] and not off[1], off)
        links_data(True); lout.clear(); lp.reload(wait_until='domcontentloaded'); lp.wait_for_selector('.ico', timeout=10000); settle(lp)
        check('P31 ...turned on, the browser asks the favicon service for those icons', len([u for u in lout if u.startswith('https://www.google.com/s2/favicons')]) == 2, lout)
        lctx.close()
        # P46: visit counts (1.53). A visitor's click counts; only the logged-in owner sees the number; Show Visit Counts off
        # hides it and the clicks still count; an encrypted tab's links are never sent.
        links_data(False)
        vctx2 = b.new_context(user_agent='Visitor/2.0', viewport={'width': 1280, 'height': 800}); vctx2.route('**/*', handle); vctx2.add_init_script(TRACK_REQUESTS)
        vp2 = vctx2.new_page(); vp2.goto(B + '/web/links/', wait_until='domcontentloaded'); vp2.wait_for_selector('a.main[data-u]', timeout=10000); settle(vp2)
        sent = []; vp2.on('request', lambda r: sent.append(r.post_data or '') if 'action=visit' in r.url else None)
        vp2.evaluate("document.querySelector('a.main[data-u]').addEventListener('click', e => e.preventDefault(), { once: true })")   # stay on the page
        vp2.click('a.main[data-u]'); until(lambda: sent, 5000); vp2.wait_for_timeout(300)
        vis_badge = vp2.evaluate("document.querySelectorAll('.vc').length")
        vctx2.close()
        octx = b.new_context(user_agent=UA, viewport={'width': 1280, 'height': 800}); octx.add_cookies(cookies); octx.route('**/*', handle); octx.add_init_script(TRACK_REQUESTS)
        op2 = octx.new_page(); op2.goto(B + '/web/links/', wait_until='domcontentloaded'); op2.wait_for_selector('a.main[data-u]', timeout=10000); settle(op2)
        shown = op2.evaluate("[...document.querySelectorAll('a.main')].map(a => (a.closest('.row').querySelector('.vc') || {}).textContent || '')")
        vpos = op2.evaluate("[...document.querySelectorAll('.row .vc')].map(v => { const r = v.closest('.row').getBoundingClientRect(), b = v.getBoundingClientRect(); return [v.parentElement.classList.contains('row'), !v.closest('a'), Math.round(r.right - b.right)] })")
        subprocess.run(['php', '-r', "require '%s/web/hub/core/lib/boot.php'; $f = hub_dir('links') . '/bookmarks.php'; $d = hub_read_guarded($f); $d['settings']['visitCounts'] = false; hub_write_guarded($f, $d);" % D], check=True)
        op2.reload(wait_until='domcontentloaded'); op2.wait_for_selector('a.main[data-u]', timeout=10000); settle(op2)
        op2.evaluate("document.querySelector('a.main[data-u]').addEventListener('click', e => e.preventDefault(), { once: true })")
        with op2.expect_response(lambda r: 'action=visit' in r.url, timeout=5000) as rv: op2.click('a.main[data-u]')
        after_off = rv.value.json()
        hidden = op2.evaluate("document.querySelectorAll('.vc').length")
        octx.close()
        check('P59 a visit count sits at the end of its row (not inside the link, beside the name), lined up with the others (1.53.8)',
              vpos and all(x[0] and x[1] and 0 <= x[2] <= 16 for x in vpos), vpos)
        check('P46 Links visit counts: a visitor\'s click counts (no number shown to them); you see it; Show Visit Counts off hides the numbers and your clicks still count',
              'example.org' in ''.join(sent) and vis_badge == 0 and shown[:1] == ['1'] and after_off.get('counted') is True and after_off.get('count') == 2 and hidden == 0, (sent[:1], vis_badge, shown, after_off, hidden))
        b.close()
finally:
    php.terminate(); shutil.rmtree(D, ignore_errors=True)
print('\n%d/%d passed' % (sum(res), len(res)))
sys.exit(0 if all(res) else 1)
