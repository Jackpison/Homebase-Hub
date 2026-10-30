#!/usr/bin/env python3
"""The Tools page (built from tools-app/), installed in a real site: the app (sections, search, favourites, recent, organise,
settings, the carry-over from the old Tools, light and dark) and every tool built so far, in a real browser, plus axe.
1.60: the Calculators, Documents and new Utilities (C1-C34), and saved presets while logged in (R1-R9).
    python3 tests/browser/tools_test.py build/site-update-X.Y.Z.zip [build/tools-new] [--browser ...]
With a second argument (python3 scripts/build-tools.py makes build/tools-new) that staging copy is swapped in first.
Needs PHP, Playwright and axe-core (npm ci). Port 8960."""
import hashlib, io, json, os, shutil, sys, zipfile

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'lib'))
from helpers import ROOT, FIX as FIXDIR, SEED
from helpers import Site, Checks, stop_all, browser_name, launch, wait_js, HERE, until, after
from playwright.sync_api import sync_playwright

BROWSER = browser_name()
if len(sys.argv) < 2: sys.exit(__doc__)
ZIP = sys.argv[1]
NEW = sys.argv[2] if len(sys.argv) > 2 else ''
DATA = json.load(open(os.path.join(ROOT, 'tools-app', 'src', 'tools.json'), encoding='utf-8'))
AXE = open(os.path.join(ROOT, 'node_modules', 'axe-core', 'axe.min.js'), encoding='utf-8').read()
RUN = "async()=>{const r=await axe.run(document.querySelector('#root'),{resultTypes:['violations'],runOnly:{type:'tag',values:['wcag2a','wcag2aa','wcag21a','wcag21aa']}});return r.violations.map(v=>v.id+' x'+v.nodes.length+' '+v.nodes[0].target.join(' '))}"
check = Checks()

try:
    s = Site(ZIP, 8960)
    t = os.path.join(s.dir, 'web', 'tools')
    if NEW and os.path.isfile(os.path.join(NEW, 'index.html')):   # a staging build: put it in place of the page
        shutil.rmtree(os.path.join(t, 'assets')); shutil.copytree(os.path.join(NEW, 'assets'), os.path.join(t, 'assets'))
        for f in ('index.html', 'sw.js'): shutil.copy(os.path.join(NEW, f), os.path.join(t, f))
    U = s.base + '/web/tools/'
    with sync_playwright() as p:
        b = launch(p, BROWSER); print('browser: %s %s' % (BROWSER, b.version))
        ctx = b.new_context(accept_downloads=True, viewport={'width': 1280, 'height': 900})
        ctx.add_init_script("if (!sessionStorage.getItem('seeded')) { localStorage.setItem('tools.favorites.v1', JSON.stringify(['image-resize', 'nope'])); localStorage.setItem('tools.recent.v1', JSON.stringify([{id: 'pdf-merge'}])); sessionStorage.setItem('seeded', '1'); }")
        pg = ctx.new_page(); errs = []
        pg.on('pageerror', lambda e: errs.append(str(e)[:200]))
        pg.goto(U); wait_js(pg, "document.querySelectorAll('#root .tc').length > 0")
        n = pg.evaluate("new Set([...document.querySelectorAll('#root .tc .tcm')].map(a => a.getAttribute('href'))).size")
        check('N1 the home lists all %d tools, with no page error' % len(DATA['tools']), n == len(DATA['tools']) and not errs, (n, errs))
        NULLS = "(() => { const w = document.createTreeWalker(document.getElementById('root'), NodeFilter.SHOW_TEXT); let n = 0; while (w.nextNode()) if (/^(null|undefined)+$/.test(w.currentNode.nodeValue.trim())) n++; return n; })()"
        check('N1b no stray "null" or "undefined" text on the home', pg.evaluate(NULLS) == 0)
        st = pg.evaluate("JSON.parse(localStorage.getItem('tools.v3'))")
        check('N2 favourites and recent tools carry over from the old Tools (unknown ids dropped)', st['fav'] == ['image-resize'] and st['recent'] == ['pdf-merge'], st)
        pg.fill('#tools-q', 'qr'); until(lambda: '#/tool/qr-generator' in pg.evaluate("[...document.querySelectorAll('#root .results .tcm')].map(a => a.getAttribute('href'))"), 3000)
        found = pg.evaluate("[...document.querySelectorAll('#root .results .tcm')].map(a => a.getAttribute('href'))")
        check('N3 searching "qr" finds the QR Code Generator', '#/tool/qr-generator' in found and len(found) < 5, found)
        pg.keyboard.press('Enter'); wait_js(pg, "location.hash === '#/tool/qr-generator' && !!document.querySelector('canvas')")
        check('N4 Enter opens the first match', pg.evaluate("location.hash") == '#/tool/qr-generator')

        def dl(click):
            with pg.expect_download(timeout=20000) as d: click()
            return open(d.value.path(), 'rb').read(), d.value.suggested_filename

        qr0 = pg.evaluate("document.querySelector('#root canvas').toDataURL()"); pg.fill('#root input[type=url]', 'https://example.com/hello')
        until(lambda: pg.evaluate("document.querySelector('#root canvas').toDataURL()") != qr0, 3000)   # the code is drawn again for the address
        data, name = dl(lambda: pg.get_by_role('button', name='Download PNG', exact=True).click())
        check('N5 QR Code Generator downloads a PNG of the code', data[:8] == b'\x89PNG\r\n\x1a\n' and name == 'qr-code.png', name)
        data, name = dl(lambda: pg.get_by_role('button', name='Download SVG', exact=True).click())
        check('N6 ...and an SVG', data.startswith(b'<svg') and b'<path' in data, data[:40])

        pg.goto(U + '#/tool/json-formatter'); wait_js(pg, "!!document.querySelector('#root textarea')")
        pg.fill('#root textarea >> nth=0', '{\n  "a": 1,\n  "b": [1, 2,]\n}'); pg.get_by_role('button', name='Format', exact=True).click(); until(lambda: pg.evaluate("!!document.querySelector('#root .note.bad')"), 3000)
        msg = pg.inner_text('#root .note.bad')
        check('N7 JSON Formatter names the exact line and column of an error', 'Line 3, column' in msg and 'comma' in msg, msg)
        pg.fill('#root textarea >> nth=0', '{"b":2,"a":{"d":1,"c":2}}'); pg.get_by_label('Sort keys A to Z').check(); pg.get_by_role('button', name='Format', exact=True).click(); until(lambda: pg.input_value('#root textarea >> nth=1').startswith('{'), 3000)
        check('N8 ...formats with sorted keys', pg.input_value('#root textarea >> nth=1') == '{\n  "a": {\n    "c": 2,\n    "d": 1\n  },\n  "b": 2\n}', pg.input_value('#root textarea >> nth=1'))

        pg.goto(U + '#/tool/hash-generator'); wait_js(pg, "!!document.querySelector('#root textarea')")
        pg.fill('#root textarea', 'abc'); until(lambda: 'ba7816bf' in pg.inner_text('#root'), 3000)
        txt = pg.inner_text('#root')
        check('N9 Hash Generator gives the known SHA-256 and SHA-1 of "abc"', 'ba7816bf8f01cfea414140de5dae2223b00361a396177a9cb410ff61f20015ad' in txt and 'a9993e364706816aba3e25717850c26c9cd0d89d' in txt)
        pg.fill('#root input[placeholder^="Paste a hash"]', 'BA7816BF8F01CFEA414140DE5DAE2223B00361A396177A9CB410FF61F20015AD'); until(lambda: 'Match:' in pg.inner_text('#root'), 3000)
        check('N10 ...and recognises a pasted hash, in capitals too', 'Match: this is the SHA-256' in pg.inner_text('#root'))

        # wait for the passwords themselves: the Hash Generator before it uses the same kind of list, so "any list" could still be its (seen once on WebKit)
        pg.goto(U + '#/tool/generators'); wait_js(pg, "(() => { const o = [...document.querySelectorAll('#root .flist .out')]; return o.length === 5 && o.every(x => x.textContent.length === 20) })()")
        pw = pg.evaluate("[...document.querySelectorAll('#root .flist .out')].map(x => x.textContent)")
        import re as _re
        check('N11 Password generator makes 5 passwords of 20 characters with every kind of character', len(pw) == 5 and all(len(x) == 20 and _re.search('[A-Z]', x) and _re.search('[a-z]', x) and _re.search('[0-9]', x) and _re.search(r'[^A-Za-z0-9]', x) for x in pw) and len(set(pw)) == 5, pw)
        pg.get_by_role('button', name='UUIDs', exact=True).click(); until(lambda: pg.evaluate("[...document.querySelectorAll('#root .flist .out')].every(x => /^[0-9a-f]{8}-/.test(x.textContent))"), 3000)
        u = pg.evaluate("[...document.querySelectorAll('#root .flist .out')].map(x => x.textContent)")
        check('N12 ...and version 4 UUIDs', len(u) == 5 and all(_re.fullmatch(r'[0-9a-f]{8}-[0-9a-f]{4}-4[0-9a-f]{3}-[89ab][0-9a-f]{3}-[0-9a-f]{12}', x) for x in u), u)

        pg.goto(U + '#/tool/text-tools'); wait_js(pg, "!!document.querySelector('#root textarea')")
        pg.fill('#root textarea', 'One two three.\n\nFour five.'); until(lambda: pg.evaluate("(document.querySelector('#root .kv div') || {}).textContent || ''").startswith('5'), 3000)
        check('N13 Word counter counts words, sentences and paragraphs', pg.evaluate("[...document.querySelectorAll('#root .kv div')].slice(0,5).map(d => d.firstChild.textContent).join(',')") == '5,26,21,2,2', pg.evaluate("[...document.querySelectorAll('#root .kv div')].map(d => d.textContent).join(' | ')"))
        pg.get_by_role('button', name='Base64', exact=True).click(); pg.fill('#root textarea >> nth=0', 'héllo ✓'); pg.get_by_role('button', name='Encode', exact=True).click(); until(lambda: pg.input_value('#root textarea >> nth=1') != '', 3000)
        enc = pg.input_value('#root textarea >> nth=1'); pg.get_by_role('button', name='Use as Input', exact=True).click(); pg.get_by_role('button', name='Decode', exact=True).click(); until(lambda: 'llo' in pg.input_value('#root textarea >> nth=1'), 3000)
        check('N14 Base64 encodes and decodes UTF-8 text', enc == 'aMOpbGxvIOKckw==' and pg.input_value('#root textarea >> nth=1') == 'héllo ✓', enc)

        pg.goto(U + '#/tool/zip-creator'); wait_js(pg, "!!document.querySelector('#root input[type=file]')")
        pg.set_input_files('#root .dropwrap input[type=file]', files=[{'name': 'a.txt', 'mimeType': 'text/plain', 'buffer': b'hello ' * 100}, {'name': 'b.txt', 'mimeType': 'text/plain', 'buffer': b'world'}])
        data, name = dl(lambda: pg.get_by_role('button', name='Create ZIP', exact=True).click())
        z = zipfile.ZipFile(io.BytesIO(data))
        check('N15 ZIP Creator puts both files into one ZIP', sorted(z.namelist()) == ['a.txt', 'b.txt'] and z.read('a.txt') == b'hello ' * 100 and name == 'files.zip', (z.namelist(), name))

        pg.goto(U + '#/tool/favicon'); wait_js(pg, "!!document.querySelector('#root .seg')")
        pg.get_by_role('button', name='Make an Icon Pack', exact=True).click(); pg.set_input_files('#root .dropwrap input[type=file]', os.path.join(FIXDIR, 'fixture-image.png'))
        wait_js(pg, "document.querySelectorAll('#root .icons img').length === 5")
        data, name = dl(lambda: pg.get_by_role('button', name='Download the Icon Pack (ZIP)', exact=True).click())
        z = zipfile.ZipFile(io.BytesIO(data)); ico = z.read('favicon.ico')
        check('N16 Favicon Maker makes a pack: favicon.ico with 3 sizes, PNGs up to 512, a manifest', {'favicon.ico', 'icon-512.png', 'apple-touch-icon.png', 'site.webmanifest'} <= set(z.namelist()) and ico[:4] == b'\x00\x00\x01\x00' and ico[4] == 3, z.namelist())

        # ---------- Images ----------
        import tempfile
        from PIL import Image as PI, ImageStat
        TMP = tempfile.mkdtemp(prefix='hub-img-')
        def img(data): return PI.open(io.BytesIO(data))
        noisy = os.path.join(TMP, 'big-photo.jpg')
        PI.effect_noise((1600, 1200), 70).convert('RGB').resize((2000, 1500)).save(noisy, quality=95)
        tif = os.path.join(TMP, 'scan.tif'); PI.open(os.path.join(FIXDIR, 'fixture-image.png')).save(tif)
        FIX = os.path.join(FIXDIR, 'fixture-image.png'); HEIC = os.path.join(FIXDIR, 'fixture-photo.heic')
        OPENED = []
        def open_tool(i):   # through the home, so the tool starts fresh even when it was the last one open
            OPENED.append(i)
            pg.goto(U + '#/'); wait_js(pg, "location.hash === '#/'"); pg.goto(U + '#/tool/' + i); wait_js(pg, "!!document.querySelector('#root .dropwrap input[type=file]')")
        def go(name, files, timeout=60000):
            after(pg, lambda: pg.set_input_files('#root .dropwrap input[type=file] >> nth=0', files))   # the tool has taken the files
            with pg.expect_download(timeout=timeout) as d: pg.get_by_role('button', name=name, exact=True).click()
            return open(d.value.path(), 'rb').read(), d.value.suggested_filename

        open_tool('image-converter')
        data, name = go('Convert', FIX); check('I1 Image Converter: PNG to JPG', data[:3] == b'\xff\xd8\xff' and name == 'fixture-image.jpg' and img(data).size == (800, 600), name)
        open_tool('image-converter'); pg.get_by_role('button', name='AVIF', exact=True).click()
        data, name = go('Convert', FIX, 120000); check('I2 ...to AVIF (the AVIF encoder loads and works)', data[4:12] == b'ftypavif' and name.endswith('.avif'), data[:16])
        open_tool('image-converter'); data, name = go('Convert', HEIC, 120000)
        check('I3 ...HEIC to JPG (the HEIC reader loads and works)', data[:3] == b'\xff\xd8\xff' and name == 'fixture-photo.jpg' and img(data).size == (320, 240), (name, data[:4]))
        open_tool('image-converter'); pg.get_by_role('button', name='PNG', exact=True).click(); data, name = go('Convert', tif)
        check('I4 ...TIFF to PNG', data[:8] == b'\x89PNG\r\n\x1a\n' and img(data).size == (800, 600), name)
        open_tool('image-converter'); after(pg, lambda: pg.set_input_files('#root .dropwrap input[type=file]', [FIX, tif]))
        pg.get_by_role('button', name='Convert', exact=True).click(); wait_js(pg, "[...document.querySelectorAll('#root button')].some(b => b.textContent.includes('Download All (ZIP)'))", 60000)
        with pg.expect_download() as d: pg.get_by_role('button', name='Download All (ZIP)').click()
        zz = zipfile.ZipFile(d.value.path()); check('I5 ...two images at once come as a ZIP of two JPGs', len(zz.namelist()) == 2 and all(n.endswith('.jpg') for n in zz.namelist()), zz.namelist())

        open_tool('image-compress'); pg.fill('#root input[aria-label="Target size"]', '120')
        data, name = go('Compress', noisy, 120000); before = os.path.getsize(noisy)
        check('I6 Image Compressor: a %d KB photo comes under the 120 KB target' % (before // 1024), len(data) <= 120 * 1024 and data[:3] == b'\xff\xd8\xff', len(data))
        check('I7 ...and the page shows the saving', 'smaller' in pg.inner_text('#root'))

        open_tool('image-resize'); pg.fill('#root input[aria-label="Width"]', '300')
        data, name = go('Resize', FIX); check('I8 Image Resizer: width 300 keeps the shape (300 x 225)', img(data).size == (300, 225) and '300x225' in name, (img(data).size, name))
        open_tool('image-resize'); pg.get_by_role('button', name='Social sizes', exact=True).click()
        data, name = go('Resize', FIX); check('I9 ...the Instagram square preset fills 1080 x 1080', img(data).size == (1080, 1080), img(data).size)

        open_tool('img-rotate'); after(pg, lambda: pg.set_input_files('#root .dropwrap input[type=file]', FIX))
        pg.get_by_role('button', name='Right', exact=True).click()
        with pg.expect_download() as d: pg.get_by_role('button', name='Save', exact=True).click()
        check('I10 Rotate and Flip: a quarter turn swaps width and height', img(open(d.value.path(), 'rb').read()).size == (600, 800))

        open_tool('img-adjust'); after(pg, lambda: pg.set_input_files('#root .dropwrap input[type=file]', FIX))
        pg.evaluate("(() => { const r = [...document.querySelectorAll('#root input[type=range]')][0]; r.value = 60; r.dispatchEvent(new Event('input')); })()")
        with pg.expect_download() as d: pg.get_by_role('button', name='Save', exact=True).click()
        lum = lambda im: sum(ImageStat.Stat(im.convert('L')).mean)
        check('I11 Image Editor: more brightness makes a brighter photo', lum(img(open(d.value.path(), 'rb').read())) > lum(PI.open(FIX)) + 20)

        open_tool('img-watermark'); data, name = go('Add to All and Save', FIX)
        check('I12 Watermark Images: same size, and the photo changed', img(data).size == (800, 600) and list(img(data).convert('RGB').get_flattened_data())[::997] != list(PI.open(FIX).convert('RGB').get_flattened_data())[::997], name)

        open_tool('img-colors'); pg.set_input_files('#root .dropwrap input[type=file]', HEIC); wait_js(pg, "document.querySelectorAll('#root .sw').length >= 4", 60000)
        hexes = pg.evaluate("[...document.querySelectorAll('#root .flist .out')].map(x => x.textContent)")
        check('I13 Color Picker: the palette finds the main colours of the photo', sum(1 for x in hexes if x.startswith('#')) >= 4, hexes[:6])

        open_tool('img-beautify'); data, name = go('Save PNG', FIX)
        check('I14 Screenshot Beautifier: padding and a window frame around the screenshot (928 x 762)', img(data).size == (928, 762), img(data).size)

        open_tool('svg-png'); pg.fill('#root textarea', '<svg xmlns="http://www.w3.org/2000/svg" width="100" height="50"><rect width="100" height="50" fill="#f00"/></svg>')
        pg.fill('#root input[type=number]', '400')
        with pg.expect_download() as d: pg.get_by_role('button', name='Convert', exact=True).click()
        sv = img(open(d.value.path(), 'rb').read()); check('I15 SVG to PNG: code at 400 pixels wide keeps the shape and colour', sv.size == (400, 200) and sv.convert('RGB').getpixel((200, 100)) == (255, 0, 0), sv.size)

        open_tool('img-base64'); pg.set_input_files('#root .dropwrap input[type=file]', FIX); until(lambda: pg.input_value('#root textarea >> nth=0').startswith('data:image'), 10000)
        uri = pg.input_value('#root textarea >> nth=0')
        check('I16 Image to Base64 makes a data URI', uri.startswith('data:image/png;base64,iVBOR'))
        pg.get_by_role('button', name='Base64 to Image', exact=True).click(); pg.fill('#root textarea', uri); pg.get_by_role('button', name='Make the File', exact=True).click()
        with pg.expect_download() as d: pg.get_by_role('button', name='Download image.png', exact=True).click()
        check('I17 ...and turns it back into the same file', open(d.value.path(), 'rb').read() == open(FIX, 'rb').read())
        # ---------- Photos ----------
        def run_save(btn='Save', timeout=60000):
            with pg.expect_download(timeout=timeout) as d: pg.get_by_role('button', name=btn, exact=True).click()
            return open(d.value.path(), 'rb').read(), d.value.suggested_filename
        def load(i, f): open_tool(i); after(pg, lambda: pg.set_input_files('#root .dropwrap input[type=file] >> nth=0', f))

        load('image-crop', FIX); pg.get_by_role('button', name='1:1', exact=True).click(); data, name = run_save('Crop and Save')
        w_, h_ = img(data).size; check('P1 Quick Crop: 1:1 gives a square', w_ == h_ and w_ >= 500, (w_, h_))
        load('image-crop', FIX); pg.get_by_role('button', name='Exact pixels', exact=True).click(); after(pg, lambda: (pg.fill('#root input[aria-label="Width in pixels"]', '300'), pg.fill('#root input[aria-label="Height in pixels"]', '200')))
        data, name = run_save('Crop and Save'); check('P2 ...exact pixels give exactly 300 x 200', img(data).size == (300, 200), img(data).size)
        pg.focus('#root .cbox'); before = pg.evaluate("document.querySelector('#root .cbox').style.left"); pg.keyboard.press('ArrowLeft'); pg.keyboard.press('ArrowLeft')
        check('P3 ...the crop box moves with the arrow keys', pg.evaluate("document.querySelector('#root .cbox').style.left") != before)

        load('circle-crop', FIX); data, name = run_save('Save PNG'); ci = img(data).convert('RGBA')
        check('P4 Circle Crop: a 512 px round PNG, see-through in the corners', ci.size == (512, 512) and ci.getpixel((2, 2))[3] == 0 and ci.getpixel((256, 256))[3] == 255, (ci.size, ci.getpixel((2, 2))))

        load('id-photo', FIX); pg.get_by_role('button', name='One photo', exact=True).click(); data, name = run_save('Make and Save')
        check('P5 Passport and ID Photo: 35 x 45 mm at 300 dpi is 413 x 531 pixels', img(data).size == (413, 531), img(data).size)
        load('id-photo', FIX); data, name = run_save('Make and Save'); txt = pg.inner_text('#root')
        check('P6 ...a 4 x 6 in print sheet (1200 x 1800) with several copies', img(data).size == (1200, 1800) and 'copies on the sheet' in txt, (img(data).size, txt[-200:]))

        load('bg-remove', FIX); wait_js(pg, "!!document.querySelector('#root .preview canvas')", 120000)
        data, name = run_save('Save'); bi = img(data).convert('RGBA'); W_, H_ = bi.size
        check('P7 Background Remover (ONNX Runtime 1.30, on the device): the background goes, the subject stays', bi.getpixel((5, 5))[3] < 80 and bi.getpixel((W_ // 2, H_ // 2))[3] > 170 and name.endswith('.png'), (bi.getpixel((5, 5)), bi.getpixel((W_ // 2, H_ // 2))))

        load('img-redact', FIX); pg.get_by_role('button', name='Add an Area in the Middle', exact=True).click(); pg.get_by_role('button', name='Solid box', exact=True).click()
        data, name = run_save(); ri = img(data).convert('RGB')
        check('P8 Blur and Redact: a solid box covers the marked area and nothing else', ri.getpixel((400, 300)) == (0, 0, 0) and ri.getpixel((10, 10)) == PI.open(FIX).convert('RGB').getpixel((10, 10)), (ri.getpixel((400, 300)), ri.getpixel((10, 10))))

        load('meme', FIX); data, name = run_save(); mi = img(data).convert('RGB'); orig = PI.open(FIX).convert('RGB')
        diff = sum(1 for x in range(0, 800, 8) if mi.getpixel((x, 60)) != orig.getpixel((x, 60)))
        check('P9 Meme and Text: the caption is drawn at the top', mi.size == (800, 600) and diff > 10, diff)

        gpsf = os.path.join(TMP, 'holiday.jpg'); ex = PI.Exif(); ex[0x010F] = 'TestMaker'; ex[0x0110] = 'Model 7'
        ex[0x8825] = {1: 'N', 2: (12.0, 58.0, 0.0), 3: 'E', 4: (77.0, 35.0, 0.0)}
        PI.new('RGB', (400, 300), (200, 100, 50)).save(gpsf, exif=ex.tobytes())
        load('image-metadata', gpsf); wait_js(pg, "document.querySelector('#root').innerText.includes('Location') || document.querySelector('#root').innerText.includes('No camera')", 20000)
        txt = pg.inner_text('#root')
        check('P10 Photo Metadata shows the camera and the GPS location, with a warning', 'TestMaker' in txt and '12.96667' in txt and 'shows where it was taken' in txt, txt[-300:])
        data, name = run_save('Remove Metadata and Save'); ci2 = img(data)
        check('P11 ...and saves the photo without any metadata', len(ci2.getexif()) == 0 and ci2.size == (400, 300), dict(ci2.getexif()))

        open_tool('collage'); cols = [(230, 60, 60), (60, 160, 230), (250, 200, 40), (60, 200, 120), (150, 80, 220), (240, 120, 30)]
        tiles = []
        for k, c in enumerate(cols): f = os.path.join(TMP, 'p%d.png' % k); PI.new('RGB', (600 + k * 50, 400 + k * 30), c).save(f); tiles.append(f)
        pg.set_input_files('#root .dropwrap input[type=file]', tiles); until(lambda: pg.evaluate("(document.querySelectorAll('#root select')[0] || {options: []}).options.length") == 21, 10000); after(pg, lambda: None, 250, 3000)
        check('P12 Collage Maker offers 21 layouts', pg.evaluate("document.querySelectorAll('#root select')[0].options.length") == 21)
        pg.select_option('#root select >> nth=0', 'g3'); data, name = run_save('Save the Collage'); co = img(data).convert('RGB')
        near = lambda p, q: all(abs(x - y) <= 6 for x, y in zip(p, q))   # JPEG rounds colours a little
        check('P13 ...a 3-column grid of 6 photos at 2400 x 2400', co.size == (2400, 2400) and near(co.getpixel((400, 600)), (230, 60, 60)) and near(co.getpixel((2000, 1800)), (240, 120, 30)), (co.size, co.getpixel((400, 600)), co.getpixel((2000, 1800))))
        pg.get_by_role('button', name='PDF', exact=True).click(); data, name = run_save('Save the Collage')
        check('P14 ...and as a PDF', data[:5] == b'%PDF-' and name == 'collage.pdf', data[:8])
        # ---------- PDF ----------
        from pypdf import PdfReader
        P3, P2 = os.path.join(FIXDIR, 'fixture-3pages.pdf'), os.path.join(FIXDIR, 'fixture-2pages.pdf')
        rd = lambda b: PdfReader(io.BytesIO(b))
        load('pdf-merge', [P3, P2]); data, name = run_save('Merge'); check('D1 Merge PDF: 3 + 2 pages make one 5-page PDF', len(rd(data).pages) == 5 and name == 'merged.pdf', name)
        load('pdf-split', P3); wait_js(pg, "document.querySelector('#root').innerText.includes('3 pages')", 20000)
        pg.get_by_role('button', name='Split', exact=True).click(); wait_js(pg, "[...document.querySelectorAll('#root button')].some(b => b.textContent.includes('Download All'))", 30000)
        with pg.expect_download() as d: pg.get_by_role('button', name='Download All (ZIP)').click()
        zz = zipfile.ZipFile(d.value.path()); check('D2 Split PDF: single pages make 3 one-page PDFs', len(zz.namelist()) == 3 and all(len(rd(zz.read(n)).pages) == 1 for n in zz.namelist()), zz.namelist())
        load('pdf-split', P3); wait_js(pg, "document.querySelector('#root').innerText.includes('3 pages')", 20000); pg.get_by_role('button', name='Your ranges', exact=True).click(); pg.fill('#root input[aria-label="Ranges"]', '1-2, 3')
        pg.get_by_role('button', name='Split', exact=True).click(); wait_js(pg, "[...document.querySelectorAll('#root button')].some(b => b.textContent.includes('Download All'))", 30000)
        with pg.expect_download() as d: pg.get_by_role('button', name='Download All (ZIP)').click()
        zz = zipfile.ZipFile(d.value.path()); check('D3 ...your ranges 1-2, 3 make files of 2 and 1 pages', sorted(len(rd(zz.read(n)).pages) for n in zz.namelist()) == [1, 2], zz.namelist())
        load('pdf-organize', P3); wait_js(pg, "document.querySelectorAll('#root .thumb').length === 3", 30000)
        pg.click('button[aria-label="Delete, page 2"]'); pg.click('button[aria-label="Rotate, page 1"]'); data, name = run_save('Save the New PDF'); r_ = rd(data)
        check('D4 Organize PDF Pages: delete one page and turn another', len(r_.pages) == 2 and r_.pages[0].rotation == 90, (len(r_.pages), r_.pages[0].rotation))
        load('pdf-to-images', P3); wait_js(pg, "document.querySelector('#root').innerText.includes('3 pages')", 20000)
        pg.get_by_role('button', name='Make Images', exact=True).click(); wait_js(pg, "[...document.querySelectorAll('#root button')].some(b => b.textContent.includes('Download All'))", 60000)
        with pg.expect_download() as d: pg.get_by_role('button', name='Download All (ZIP)').click()
        zz = zipfile.ZipFile(d.value.path()); check('D5 PDF to Images: 3 pages make 3 JPGs', len(zz.namelist()) == 3 and all(zz.read(n)[:3] == b'\xff\xd8\xff' for n in zz.namelist()), zz.namelist())
        load('images-to-pdf', [FIX, HEIC]); data, name = run_save('Make PDF'); r_ = rd(data)
        dims = [sorted((round(float(p.mediabox.width)), round(float(p.mediabox.height)))) for p in r_.pages]
        check('D6 Images to PDF: a PNG and a HEIC make a 2-page A4 PDF (landscape photos on landscape pages)', len(r_.pages) == 2 and dims == [[595, 842], [595, 842]] and float(r_.pages[0].mediabox.width) > float(r_.pages[0].mediabox.height), dims)
        load('pdf-compress', P3); pg.get_by_role('button', name='Light (keeps the text)', exact=True).click(); data, name = run_save('Compress', 60000)
        check('D7 PDF Compressor: Light keeps a valid PDF with its 3 pages and its text', len(rd(data).pages) == 3 and 'quick brown fox' in rd(data).pages[0].extract_text().lower(), name)
        load('pdf-compress', P3); data, name = run_save('Compress', 120000); check('D8 ...Strong rebuilds the pages', len(rd(data).pages) == 3, name)
        load('pdf-protect', P3); pg.fill('#root input[type=password] >> nth=0', 'open sesame'); pg.fill('#root input[type=password] >> nth=1', 'open sesame'); data, name = run_save('Protect', 60000)
        r_ = rd(data); enc = r_.is_encrypted; ed = str(r_.trailer['/Encrypt'].get_object()) if enc else ''   # read without decrypting: CI has no cryptography package
        check('D9 Protect or Unlock: AES-256 with a password (qpdf)', enc and '/AESV3' in ed and "'/V': 5" in ed, ed[:160])
        prot = os.path.join(TMP, 'locked.pdf'); open(prot, 'wb').write(data)
        open_tool('pdf-protect'); pg.get_by_role('button', name='Remove a Password', exact=True).click(); after(pg, lambda: pg.set_input_files('#root .dropwrap input[type=file]', prot))
        pg.fill('#root input[type=password]', 'wrong one'); pg.get_by_role('button', name='Remove the Password', exact=True).click()
        wait_js(pg, "document.body.innerText.includes('not right')", 30000)
        check('D10 ...a wrong password is refused with a plain message', 'That password is not right.' in pg.inner_text('body'))
        pg.fill('#root input[type=password]', 'open sesame'); data, name = run_save('Remove the Password', 60000)
        check('D11 ...and the right one removes it', not rd(data).is_encrypted and len(rd(data).pages) == 3, name)
        load('pdf-text', P3); pg.get_by_role('button', name='Text Layer Only', exact=True).click(); pg.get_by_role('button', name='Get the Text', exact=True).click()
        wait_js(pg, "document.querySelector('#root textarea').value.length > 10", 60000); txt = pg.input_value('#root textarea')
        check('D12 PDF to Text reads the text layer, page by page', 'quick brown fox' in txt.lower() and '--- page 3 ---' in txt, txt[:120])
        load('pdf-text', os.path.join(FIXDIR, 'japanese-cid.pdf')); pg.get_by_role('button', name='Text Layer Only', exact=True).click(); pg.get_by_role('button', name='Get the Text', exact=True).click()
        wait_js(pg, "document.querySelector('#root textarea').value.length > 10", 60000); txt = pg.input_value('#root textarea')
        check('D12b ...and Japanese text in a PDF that does not carry its font (pdf.js character maps)', '日本語のテキスト' in txt, txt[:120])
        ocr = os.path.join(TMP, 'note.png'); im_ = PI.new('RGB', (900, 260), 'white'); from PIL import ImageDraw, ImageFont
        ImageDraw.Draw(im_).text((40, 70), 'HELLO WORLD 2026', fill='black', font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf', 72)); im_.save(ocr)
        load('pdf-text', ocr); pg.get_by_role('button', name='Get the Text', exact=True).click()
        wait_js(pg, "document.querySelector('#root textarea').value.length > 5", 180000); txt = pg.input_value('#root textarea')
        check('D13 ...and reads a photo of text with OCR (Tesseract 7, in the browser)', 'HELLO WORLD' in txt.upper(), txt[:120])
        load('pdf-stamp', P3); data, name = run_save('Stamp and Save', 60000); t1 = rd(data).pages[0].extract_text()
        check('D14 Watermark and Page Numbers: CONFIDENTIAL and "Page 1 of 3"', 'CONFIDENTIAL' in t1 and 'Page 1 of 3' in t1 and 'Page 3 of 3' in rd(data).pages[2].extract_text(), t1[-120:])
        load('pdf-sign', P3); wait_js(pg, "!!document.querySelector('#root .stage canvas')", 30000)
        pg.locator('#root canvas.pad').scroll_into_view_if_needed()
        box = pg.locator('#root canvas.pad').bounding_box(); W2, H2 = box['width'], box['height']   # draw a zigzag across the pad, whatever its size
        pg.mouse.move(box['x'] + W2 * 0.1, box['y'] + H2 * 0.5); pg.mouse.down()
        for k in range(1, 12): pg.mouse.move(box['x'] + W2 * (0.1 + k * 0.07), box['y'] + H2 * (0.5 + (0.25 if k % 2 else -0.25)))
        pg.mouse.up(); until(lambda: pg.evaluate("!!document.querySelector('#root .sigprev')"), 3000)   # the signature is taken (its preview shows)
        pg.click('#root .layer', position={'x': 300, 'y': 500}); until(lambda: pg.evaluate("document.querySelectorAll('#root .layer .item').length") == 1, 3000)
        check('D15 Sign PDF: a drawn signature is placed on the page', pg.evaluate("document.querySelectorAll('#root .layer .item').length") == 1)
        data, name = run_save('Save the Signed PDF', 60000); res = rd(data).pages[0]['/Resources']
        check('D16 ...and saved into the PDF as an image', '/XObject' in res and len(res['/XObject']) >= 1, list(res.keys()))
        load('pdf-annotate', P3); wait_js(pg, "!!document.querySelector('#root .stage canvas')", 30000); pg.fill('#root input[aria-label="Text to add"]', 'Hello PDF')
        pg.click('#root .layer', position={'x': 200, 'y': 300}); pg.get_by_role('button', name='Tick', exact=True).click(); pg.click('#root .layer', position={'x': 200, 'y': 400})
        pg.get_by_role('button', name='Highlight', exact=True).click(); pg.click('#root .layer', position={'x': 400, 'y': 200}); until(lambda: pg.evaluate("document.querySelectorAll('#root .layer .item').length") == 3, 3000)
        check('D17 Fill and Annotate: text, a tick and a highlight are placed', pg.evaluate("document.querySelectorAll('#root .layer .item').length") == 3)
        data, name = run_save('Save the PDF', 60000); check('D18 ...and the text is in the saved PDF', 'Hello PDF' in rd(data).pages[0].extract_text(), name)
        shutil.rmtree(TMP, ignore_errors=True)

        # the app itself
        pg.goto(U + '#/'); wait_js(pg, "document.querySelectorAll('#root .tc').length > 0")
        pg.click('#root .tc:has(a[href="#/tool/json-formatter"]) .star'); pg.goto(U + '#/favourites'); wait_js(pg, "location.hash === '#/favourites'"); until(lambda: '#/tool/json-formatter' in pg.evaluate("[...document.querySelectorAll('#root .tcm')].map(a => a.getAttribute('href'))"), 3000)
        favs = pg.evaluate("[...document.querySelectorAll('#root .tcm')].map(a => a.getAttribute('href'))")
        check('N17 a star adds a favourite, and Favourites lists it', '#/tool/json-formatter' in favs and '#/tool/image-resize' in favs, favs)
        pg.goto(U + '#/recent'); until(lambda: len(pg.evaluate("[...document.querySelectorAll('#root .tcm')].map(a => a.getAttribute('href'))")) >= 3, 3000)
        rec = pg.evaluate("[...document.querySelectorAll('#root .tcm')].map(a => a.getAttribute('href'))")
        last = []
        for i in reversed(OPENED):
            if i not in last: last.append(i)
        check('N18 Recent lists the tools just used, newest first, and keeps at most 12', rec[:3] == ['#/tool/' + i for i in last[:3]] and len(rec) == 12, (rec[:4], last[:3], len(rec)))
        pg.goto(U + '#/organise'); pg.wait_for_selector('button[aria-label="Show QR Code Generator"]', timeout=5000)
        pg.click('button[aria-label="Show QR Code Generator"]'); pg.goto(U + '#/s/utilities'); until(lambda: '#/tool/favicon' in pg.evaluate("[...document.querySelectorAll('#root .tcm')].map(a => a.getAttribute('href'))"), 3000)
        ul = pg.evaluate("[...document.querySelectorAll('#root .tcm')].map(a => a.getAttribute('href'))")
        check('N19 Organise hides a tool from its section', '#/tool/qr-generator' not in ul and '#/tool/favicon' in ul, ul)
        pg.goto(U + '#/organise'); pg.wait_for_selector('button[aria-label^="Move Favicon"]', timeout=5000); pg.click('button[aria-label="Move Favicon Downloader down"], button[aria-label^="Move Favicon"][aria-label$="down"]')
        pg.goto(U + '#/s/utilities'); until(lambda: '#/tool/favicon' in pg.evaluate("[...document.querySelectorAll('#root .tcm')].map(a => a.getAttribute('href'))") and pg.evaluate("[...document.querySelectorAll('#root .tcm')].map(a => a.getAttribute('href'))") != ul, 3000)
        ul2 = pg.evaluate("[...document.querySelectorAll('#root .tcm')].map(a => a.getAttribute('href'))")
        check('N20 ...and changes the order within a section', ul2.index('#/tool/favicon') > ul.index('#/tool/favicon'), (ul, ul2))
        pg.goto(U + '#/settings'); pg.get_by_role('button', name='Tabs on Top', exact=True).wait_for(timeout=5000); pg.get_by_role('button', name='Tabs on Top', exact=True).click(); pg.reload(); wait_js(pg, "!!document.querySelector('.tl')")
        check('N21 Settings: tabs on top is kept after a reload, with no stray text in the menu', pg.evaluate("document.querySelector('.tl').dataset.layout") == 'tabs' and pg.evaluate(NULLS) == 0)
        dark = pg.evaluate("getComputedStyle(document.body).backgroundColor"); pg.evaluate("document.documentElement.setAttribute('data-theme','light')"); until(lambda: pg.evaluate("getComputedStyle(document.body).backgroundColor") != dark, 3000)
        check('N22 light and dark follow the site', dark != pg.evaluate("getComputedStyle(document.body).backgroundColor"))
        missing = [t['id'] for t in DATA['tools']]
        pg.goto(U + '#/tool/nothing-like-this'); until(lambda: 'Tool not found' in pg.inner_text('#root'), 5000)
        check('N23 an unknown tool address says so and offers the list', 'Tool not found' in pg.inner_text('#root'))
        # ---------- 1.60: Calculators, Documents and the new Utilities ----------
        TMP2 = tempfile.mkdtemp(prefix='hub-new-')
        def tool(i): pg.goto(U + '#/'); wait_js(pg, "location.hash === '#/'"); pg.goto(U + '#/tool/' + i); wait_js(pg, "!!document.querySelector('#root .th h1') && !document.querySelector('#root [aria-busy=true]')", 20000)
        def txt(): return pg.inner_text('#root .view')
        def stat(label): return pg.evaluate("(l) => { const s = [...document.querySelectorAll('#root .stat')].find(x => x.querySelector('span').textContent === l); return s ? s.querySelector('b').textContent : null }", label)
        def setf(label, v): pg.get_by_label(label, exact=True).fill(str(v))
        tool('loan-emi'); setf('Loan Amount', 1000000); setf('Interest Rate', 10); setf('Loan Term', 10)
        until(lambda: stat('Monthly EMI') == '₹13,215', 3000)
        check('C1 EMI Calculator: 10 lakh at 10% for 10 years is ₹13,215 a month (₹5,85,809 interest)', stat('Monthly EMI') == '₹13,215' and stat('Total Interest') == '₹5,85,809', (stat('Monthly EMI'), stat('Total Interest')))
        rows = pg.evaluate("document.querySelectorAll('#root .tbl tbody tr').length"); last = pg.evaluate("[...document.querySelectorAll('#root .tbl tbody tr:last-child td')].map(t => t.textContent)")
        check('C2 ...the schedule has 10 years and ends at a balance of ₹0', rows == 10 and last[-1] == '₹0', (rows, last))
        data, name = dl(lambda: pg.get_by_role('button', name='Download CSV', exact=True).click()); lines = data.decode('utf-8-sig').splitlines()
        check('C3 ...and downloads all 120 months as a CSV', name == 'loan-schedule.csv' and len(lines) == 121 and lines[0] == 'Month,Principal,Interest,Balance', (name, len(lines), lines[:2]))
        tool('sip-calculator'); setf('Monthly Investment', 10000); setf('Expected Return', 12); setf('Time', 10)
        until(lambda: stat('Total Value') == '₹23,23,391', 3000)
        check('C4 SIP: ₹10,000 a month at 12% for 10 years grows to ₹23,23,391 (₹12,00,000 invested)', stat('Total Value') == '₹23,23,391' and stat('Invested') == '₹12,00,000', (stat('Total Value'), stat('Invested')))
        setf('Step Up Each Year', 10); until(lambda: stat('Invested') == '₹19,12,491', 3000)
        check('C5 ...with a 10% step-up each year, ₹19,12,491 goes in', stat('Invested') == '₹19,12,491', stat('Invested'))
        pg.get_by_role('button', name='Lumpsum (One Time)', exact=True).click(); setf('Amount Invested', 100000); until(lambda: stat('Total Value') == '₹3,10,585', 3000)
        check('C6 ...a lumpsum of ₹1,00,000 at 12% for 10 years is ₹3,10,585', stat('Total Value') == '₹3,10,585', stat('Total Value'))
        tool('fd-rd-calculator'); setf('Deposit', 100000); setf('Interest Rate', 7); setf('Years', 5)
        until(lambda: stat('Maturity Value') == '₹1,41,478', 3000)
        check('C7 FD: ₹1,00,000 at 7% for 5 years, added every quarter, matures at ₹1,41,478', stat('Maturity Value') == '₹1,41,478', stat('Maturity Value'))
        pg.get_by_role('button', name='Recurring Deposit (RD)', exact=True).click(); setf('Monthly Deposit', 5000); setf('Years', 1); setf('Interest Rate', 7)
        until(lambda: stat('Maturity Value') == '₹62,311', 3000)
        check('C8 RD: ₹5,000 a month at 7% for 12 months matures at ₹62,311 (quarterly compounding, as Indian banks count it)', stat('Maturity Value') == '₹62,311', stat('Maturity Value'))
        tool('percentage-calculator')
        outs = pg.evaluate("[...document.querySelectorAll('#root .big-out')].map(x => x.textContent)")
        check('C9 Percentage: 15% of 2400 is 360, 300 of 2400 is 12.5%, 1200 to 1500 is +25%, 20% off 2499 is 1,999.2', outs[:4] == ['360', '12.5%', '+25%', '1,999.2'] and outs[4] == '720' and outs[6] == '1,000', outs)
        tool('age-date-calculator'); pg.get_by_label('Date of Birth', exact=True).fill('2000-02-29'); pg.get_by_label('Age On', exact=True).fill('2026-10-27')
        until(lambda: '26 years, 7 months, 28 days' in txt(), 3000)
        check('C10 Age: born 29 Feb 2000, on 27 Oct 2026 the age is 26 years, 7 months, 28 days', '26 years, 7 months, 28 days' in txt(), txt()[:300])
        pg.get_by_role('button', name='Between Two Dates', exact=True).click(); pg.get_by_label('From', exact=True).fill('2026-10-01'); pg.get_by_label('To', exact=True).fill('2026-10-31')
        until(lambda: '30 days' in txt(), 3000)
        check('C11 ...1 to 31 October 2026 is 30 days, with 22 working days (Mon–Fri)', '30 days' in txt() and stat('Working Days (Mon–Fri)') == '22', (stat('Working Days (Mon–Fri)'),))
        pg.get_by_role('button', name='Add or Subtract', exact=True).click(); pg.get_by_label('Start Date', exact=True).fill('2026-01-31'); pg.get_by_label('Of', exact=True).select_option('m'); setf('How Many', 1)
        until(lambda: '2026-02-28' in txt(), 3000)
        check('C12 ...31 January plus 1 month is 28 February (the end of the month, not March)', '2026-02-28' in txt(), txt()[:200])
        tool('unit-converter')
        pg.get_by_role('button', name='Temperature', exact=True).click(); setf('Value', 100); until(lambda: '212' in pg.inner_text('#root .big-out'), 3000)
        t1 = pg.inner_text('#root .big-out'); pg.get_by_role('button', name='Area', exact=True).click(); setf('Value', 1); pg.get_by_label('From', exact=True).select_option('ac'); pg.get_by_label('To', exact=True).select_option('guntha')
        until(lambda: pg.inner_text('#root .big-out').startswith('40 '), 3000)
        check('C13 Unit Converter: 100 °C is 212 °F, and 1 acre is 40 guntha', t1.startswith('212') and pg.inner_text('#root .big-out').startswith('40 '), (t1, pg.inner_text('#root .big-out')))
        tool('number-to-words'); pg.get_by_label('Number', exact=True).fill('1234567.89'); until(lambda: 'Lakh' in pg.inner_text('#root .words'), 3000)
        w1 = pg.inner_text('#root .words'); pg.get_by_role('button', name='International (Million)', exact=True).click(); pg.get_by_label('As money (for a cheque or invoice)').uncheck(); pg.get_by_label('Number', exact=True).fill('1000000001')
        until(lambda: 'Billion' in pg.inner_text('#root .words'), 3000); w2 = pg.inner_text('#root .words')
        check('C14 Number to Words: 1234567.89 is "Twelve Lakh Thirty-Four Thousand Five Hundred Sixty-Seven Rupees and Eighty-Nine Paise Only"', w1 == 'Twelve Lakh Thirty-Four Thousand Five Hundred Sixty-Seven Rupees and Eighty-Nine Paise Only', w1)
        check('C15 ...and 1000000001 the international way is "One Billion One"', w2 == 'One Billion One', w2)

        # Documents
        tool('text-compare'); pg.get_by_label('Original text').fill('one\ntwo\nthree\nfour'); pg.get_by_label('Changed text').fill('one\n2\nthree\nfour\nfive')
        until(lambda: '2 lines added, 1 removed' in txt(), 3000)
        check('C16 Text Compare: line by line it finds 2 added and 1 removed', '2 lines added, 1 removed' in txt() and pg.evaluate("document.querySelectorAll('#root .diff .del').length") == 1, txt()[-200:])
        pg.get_by_role('button', name='Word by Word', exact=True).click(); pg.get_by_label('Original text').fill('the quick brown fox'); pg.get_by_label('Changed text').fill('the slow brown fox')
        until(lambda: pg.evaluate("[...document.querySelectorAll('#root del.dw')].map(x => x.textContent).join()") == 'quick', 3000)
        check('C17 ...word by word it marks "quick" removed and "slow" added', pg.evaluate("[...document.querySelectorAll('#root ins.dw')].map(x => x.textContent).join()") == 'slow')
        tool('markdown-editor'); pg.get_by_label('Markdown text').fill('# Hi\n\n**bold** [ok](https://example.com) [bad](javascript:alert(1)) ![p](https://example.com/x.png) <b>raw</b>\n\n- [x] done\n\n| a | b |\n|---|--:|\n| 1 | 2 |')
        until(lambda: pg.evaluate("(document.querySelector('#root .md h1') || {}).textContent === 'Hi'"), 3000)
        md = pg.evaluate("(() => { const m = document.querySelector('#root .md'); return { h1: m.querySelector('h1')?.textContent, b: m.querySelector('strong')?.textContent, links: [...m.querySelectorAll('a')].map(a => a.getAttribute('href')), imgs: m.querySelectorAll('img').length, raw: m.querySelectorAll('b').length, box: !!m.querySelector('input[type=checkbox]:checked') } })()")
        check('C18 Markdown Editor: headings, bold, a task and a table show; only the https link is a link; no picture or raw HTML is loaded', md['h1'] == 'Hi' and md['b'] == 'bold' and md['links'] == ['https://example.com', 'https://example.com/x.png'] and md['imgs'] == 0 and md['raw'] == 0 and md['box'], md)
        data, name = dl(lambda: pg.get_by_role('button', name='Save .html', exact=True).click())
        check('C19 ...saves as HTML named after its heading', name == 'hi.html' and b'<h1>Hi</h1>' in data and b'javascript:' not in data, name)
        pg.reload(); wait_js(pg, "!!document.querySelector('#root .md h1')", 10000)
        check('C20 ...and keeps the draft after a reload', pg.input_value('#root textarea').startswith('# Hi'))
        tool('csv-json-converter'); pg.get_by_role('button', name='JSON', exact=True).click()
        until(lambda: '"city": "Delhi, NCR"' in pg.input_value('#root textarea[aria-label="Converted text"]'), 3000)
        js = json.loads(pg.input_value('#root textarea[aria-label="Converted text"]'))
        check('C21 CSV to JSON: quoted commas kept, numbers typed', js == [{'name': 'Asha', 'city': 'Pune', 'amount': 1200}, {'name': 'Ravi', 'city': 'Delhi, NCR', 'amount': 850.5}], js)
        pg.get_by_role('button', name='Excel (.xlsx)', exact=True).click(); data, name = dl(lambda: pg.get_by_role('button', name='Download', exact=True).click())
        import openpyxl
        xl = os.path.join(TMP2, 'data.xlsx'); open(xl, 'wb').write(data); ws = openpyxl.load_workbook(xl).active
        check('C22 ...to Excel: Excel\'s format, read back with openpyxl (header bold, numbers as numbers)', name == 'data.xlsx' and ws['B3'].value == 'Delhi, NCR' and ws['C3'].value == 850.5 and ws['A1'].font.b, (name, [[c.value for c in r] for r in ws.iter_rows()]))
        wb = openpyxl.Workbook(); s1 = wb.active; s1.title = 'Sales'; s1.append(['Item', 'Qty', 'When']); import datetime; s1.append(['Tea', 3, datetime.date(2026, 10, 1)]); wb.create_sheet('Other').append(['x']); xin = os.path.join(TMP2, 'in.xlsx'); wb.save(xin)
        pg.get_by_role('button', name='CSV', exact=True).click(); after(pg, lambda: pg.set_input_files('#root .dropwrap input[type=file]', xin))
        until(lambda: 'Tea,3,2026-10-01' in pg.input_value('#root textarea[aria-label="Converted text"]'), 5000)
        check('C23 ...an Excel file (two sheets, a date) back to CSV', 'Item,Qty,When' in pg.input_value('#root textarea[aria-label="Converted text"]') and 'Tea,3,2026-10-01' in pg.input_value('#root textarea[aria-label="Converted text"]') and pg.is_visible('#root select'), pg.input_value('#root textarea[aria-label="Converted text"]'))
        pg.evaluate("localStorage.removeItem('tools.invoice.v1')"); tool('invoice-generator')
        setf('Business Name', 'Sharma Traders'); setf('Customer Name', 'Mehta & Sons'); setf('Description 1', 'Steel rods, 12 mm'); setf('Quantity 1', '10'); setf('Rate 1', '450')
        pg.get_by_role('button', name='Add Item', exact=True).click(); setf('Description 2', 'Delivery'); setf('Rate 2', '500'); pg.get_by_label('Split the tax into two halves').check()
        until(lambda: pg.inner_text('#root .total') == '₹5,900.00', 3000)
        check('C24 Invoice Generator: 10 × 450 + 500 with 18% tax split into CGST and SGST totals ₹5,900.00', pg.inner_text('#root .total') == '₹5,900.00' and 'CGST (9%)' in txt(), pg.inner_text('#root .total'))
        data, name = dl(lambda: pg.get_by_role('button', name='Download PDF', exact=True).click())
        from pypdf import PdfReader
        pt = PdfReader(io.BytesIO(data)).pages[0].extract_text()
        check('C25 ...the PDF has the business, customer, items, CGST and SGST at 9%, the total and the total in words', name == 'invoice-INV-001.pdf' and all(x in pt for x in ['Sharma Traders', 'Mehta & Sons', 'Steel rods, 12 mm', 'CGST @ 9%', 'SGST @ 9%', '5,900.00', 'Five Thousand Nine Hundred Rupees Only']), pt[:600])
        open(os.path.join(TMP2, 'invoice.pdf'), 'wb').write(data)

        # Utilities
        pg.goto(U + '#/tool/qr-generator'); wait_js(pg, "!!document.querySelector('#root canvas')"); q0 = pg.evaluate("document.querySelector('#root canvas').toDataURL()")
        pg.fill('#root input[type=url]', 'https://example.com/scan-me'); until(lambda: pg.evaluate("document.querySelector('#root canvas').toDataURL()") != q0, 3000)
        data, name = dl(lambda: pg.get_by_role('button', name='Download PNG', exact=True).click()); qp = os.path.join(TMP2, 'qr.png'); open(qp, 'wb').write(data)
        tool('qr-scanner'); wasm = []
        pg.on('request', lambda r: wasm.append(r.url) if '.wasm' in r.url else None)
        after(pg, lambda: pg.set_input_files('#root .dropwrap input[type=file]', qp)); until(lambda: 'https://example.com/scan-me' in txt(), 30000)
        check('C26 QR and Barcode Scanner: reads a QR code from a picture (ZXing from this site, not from a CDN)', 'https://example.com/scan-me' in txt() and all(u.startswith(s.base) for u in wasm), (txt()[-300:], wasm))
        tool('barcode-generator'); until(lambda: pg.evaluate("!!document.querySelector('#root .bcprev svg rect')"), 5000)
        data, name = dl(lambda: pg.get_by_role('button', name='Download PNG', exact=True).click())
        check('C27 Barcode Generator: a Code 128 barcode as PNG', data[:8] == b'\x89PNG\r\n\x1a\n' and name == 'barcode-Homebase-001.png', name)
        bp = os.path.join(TMP2, 'bc.png'); open(bp, 'wb').write(data)
        pg.get_by_label('Type', exact=True).select_option('EAN13'); until(lambda: pg.evaluate("!!document.querySelector('#root .bcprev svg')"), 3000)
        pg.get_by_label('Value', exact=True).fill('12345'); until(lambda: pg.evaluate("!document.querySelector('#root .bcprev svg')"), 3000)
        check('C28 ...an EAN-13 needs 12 or 13 digits: a short number is refused with a reason', 'cannot be made' in txt() and pg.get_by_role('button', name='Download PNG', exact=True).is_disabled())
        tool('qr-scanner'); after(pg, lambda: pg.set_input_files('#root .dropwrap input[type=file]', bp)); until(lambda: 'Homebase-001' in txt(), 30000)
        check('C29 ...and the scanner reads that barcode back', 'Homebase-001' in txt() and 'Code 128' in txt(), txt()[-200:])
        pg.evaluate("localStorage.setItem('tools.timezones.v1', JSON.stringify(['Asia/Kolkata', 'America/New_York', 'Asia/Tokyo']))"); tool('timezone-converter')
        pg.get_by_label('In This Time Zone', exact=True).fill('Asia/Kolkata'); pg.get_by_label('In This Time Zone', exact=True).dispatch_event('change'); pg.get_by_label('Date and Time', exact=True).fill('2026-07-01T19:30')
        until(lambda: 'next day' in txt(), 3000); rws = pg.evaluate("[...document.querySelectorAll('#root .tzrow b')].map(b => b.textContent)")
        check('C30 Time Zone Converter: 7:30 pm in India on 1 July is 10:00 am in New York (summer time) and 11:00 pm in Tokyo', [r.replace(' ', ' ').lower() for r in rws] == ['7:30 pm', '10:00 am', '11:00 pm'], rws)
        tool('stopwatch-timer'); pg.get_by_role('button', name='Start', exact=True).click(); pg.wait_for_timeout(1300); pg.get_by_role('button', name='Lap', exact=True).click(); pg.get_by_role('button', name='Pause', exact=True).click()
        c1 = pg.inner_text('#root .clock')
        check('C31 Stopwatch: runs, takes a lap and pauses', c1.startswith('00:01') and pg.evaluate("document.querySelectorAll('#root .tbl tbody tr').length") == 1, c1)
        pg.get_by_role('button', name='Timer', exact=True).click(); pg.get_by_label('Minutes', exact=True).fill('0'); pg.get_by_label('Seconds', exact=True).fill('2'); pg.get_by_role('button', name='Start', exact=True).click()
        until(lambda: 'Time is up.' in txt(), 6000)
        check('C32 ...the timer counts down and says when time is up', 'Time is up.' in txt() and pg.inner_text('#root .clock') == '00:00')
        tool('random-picker'); pg.get_by_role('button', name='Pick', exact=True).click()
        got = pg.evaluate("[...document.querySelectorAll('#root .picked b')].map(b => b.textContent)")
        pg.get_by_role('button', name='Numbers', exact=True).click(); setf('To', 10); setf('How Many', 10); pg.get_by_role('button', name='Pick', exact=True).click()
        nums = pg.evaluate("[...document.querySelectorAll('#root .picked b')].map(b => +b.textContent)")
        check('C33 Random Picker: picks one name from the list, and 10 different numbers from 1 to 10', len(got) == 1 and got[0] in ['Asha', 'Ravi', 'Meera', 'Kabir', 'Zoya', 'Arjun'] and sorted(nums) == list(range(1, 11)), (got, nums))
        tool('recorder')
        check('C34 Screen and Voice Recorder: opens on Screen on a computer, with Start Recording', pg.get_by_role('button', name='Start Recording', exact=True).is_visible() and pg.get_by_role('button', name='Screen', exact=True).get_attribute('aria-pressed') == 'true')
        shutil.rmtree(TMP2, ignore_errors=True)
        # N26 (1.60): every tool in tools.json opens: its heading shows and nothing fails (a tool added with scripts/new-tool.py is covered at once)
        broken = []
        for t in DATA['tools']:
            pg.goto(U + '#/tool/' + t['id'])
            ok = until(lambda: pg.evaluate("(() => { const h = document.querySelector('#root .th h1'), b = document.querySelector('#root .th ~ *'); return !!h && h.textContent.trim() === %r && !document.querySelector('#root [aria-busy=true]') })()" % t['name']), 20000, 150)
            if not ok or pg.evaluate("!!document.querySelector('#root .note.bad[role=alert]')"): broken.append(t['id'])
        check('N26 every tool (%d) opens with its heading and without an error' % len(DATA['tools']), not broken, broken)
        check('N24 no page error anywhere', not errs, errs)
        for scheme in ('dark', 'light'):
            pg.evaluate("document.documentElement.setAttribute('data-theme', %r)" % scheme)
            found = []
            for u in ('#/', '#/tool/qr-generator', '#/tool/json-formatter', '#/tool/generators', '#/tool/image-converter', '#/tool/image-resize', '#/tool/img-adjust', '#/tool/img-watermark', '#/tool/collage', '#/tool/id-photo', '#/tool/image-crop', '#/tool/img-redact', '#/tool/meme', '#/tool/bg-remove', '#/tool/pdf-merge', '#/tool/pdf-split', '#/tool/pdf-protect', '#/tool/pdf-text', '#/tool/pdf-stamp', '#/tool/pdf-sign', '#/tool/pdf-annotate', '#/tool/loan-emi', '#/tool/sip-calculator', '#/tool/age-date-calculator', '#/tool/invoice-generator', '#/tool/text-compare', '#/tool/markdown-editor', '#/tool/csv-json-converter', '#/tool/timezone-converter', '#/tool/random-picker', '#/tool/barcode-generator', '#/organise', '#/settings'):
                after(pg, lambda: pg.goto(U + u), 250, 10000); pg.evaluate("document.documentElement.setAttribute('data-theme', %r)" % scheme); pg.mouse.move(2, 2)
                pg.wait_for_timeout(100)   # on purpose: a hovered button fades back (a CSS transition) before axe reads the colours
                pg.evaluate(AXE); v = pg.evaluate(RUN)
                if v: found.append(u + ': ' + '; '.join(v[:3]))
            check('N25 axe finds nothing on the home, the tools named above, Organise and Settings (%s)' % scheme, not found, found[:4])
        # ---------- 1.60: saved presets (on the server, only while logged in) ----------
        import urllib.request, urllib.error
        from helpers import browser_login
        try: code = urllib.request.urlopen(s.base + '/web/hub/core/api.php?action=presets').status
        except urllib.error.HTTPError as e: code = e.code
        pc = b.new_context(accept_downloads=True, viewport={'width': 1280, 'height': 900}); pp = pc.new_page(); perr = []
        pp.on('pageerror', lambda e: perr.append(str(e)[:200])); pp.on('dialog', lambda d: d.accept())
        pp.goto(U + '#/tool/qr-generator'); until(lambda: 'Log in to save' in (pp.evaluate("(document.querySelector('#root .pbar') || {}).textContent || ''")), 8000)
        check('R1 logged out: a tool offers to log in to save presets, and the server gives nothing (401)', 'Log in to save' in pp.inner_text('#root .pbar') and code == 401, code)
        check('R2 logging in works from Tools', browser_login(pp, s.base) == '')
        pp.goto(U + '#/tool/qr-generator'); until(lambda: pp.evaluate("!!document.querySelector('#root .pbar select')"), 8000)
        pp.fill('#root input[type=color] >> nth=0', '#1a3fa0'); pp.get_by_role('button', name='Save as Preset', exact=True).click()
        pp.get_by_label('Preset name').fill('Blue Code'); pp.get_by_role('button', name='Save', exact=True).click()
        until(lambda: 'Blue Code' in pp.evaluate("[...document.querySelectorAll('#root .pbar option')].map(o => o.textContent).join()"), 5000)
        pp.reload(); until(lambda: 'Blue Code' in pp.evaluate("[...document.querySelectorAll('#root .pbar option')].map(o => o.textContent).join()"), 8000)
        before = pp.input_value('#root input[type=color] >> nth=0'); pp.select_option('#root .pbar select', label='Blue Code')
        until(lambda: pp.input_value('#root input[type=color] >> nth=0') == '#1a3fa0', 3000)
        check('R3 QR Code Generator: Save as Preset keeps the colours on the server; after a reload, choosing it brings them back', before == '#000000' and pp.input_value('#root input[type=color] >> nth=0') == '#1a3fa0', before)
        pp.goto(U + '#/tool/pdf-sign'); until(lambda: pp.evaluate("!!document.querySelector('#root .pbar select')"), 8000)
        pp.get_by_role('button', name='Type', exact=True).click(); pp.get_by_label('Your name', exact=True).fill('Asha Rao')
        pp.get_by_role('button', name='Save as Preset', exact=True).click(); pp.get_by_label('Preset name').fill('My Signature'); pp.get_by_role('button', name='Save', exact=True).click()
        until(lambda: 'sign' in (s.get('presets').get('kinds') or {}), 5000)
        sg = (s.get('presets').get('kinds') or {}).get('sign', [{}])[0]
        check('R4 Sign PDF: the signature is saved as a preset (a PNG picture)', sg.get('name') == 'My Signature' and sg.get('data', {}).get('img', '').startswith('data:image/png;base64,'), str(sg)[:200])
        pp.evaluate("localStorage.removeItem('tools.invoice.v1')"); pp.goto(U + '#/tool/invoice-generator'); until(lambda: pp.evaluate("!!document.querySelector('#root .pbar select')"), 8000)
        pp.get_by_label('Business Name', exact=True).fill('Sharma Traders'); pp.get_by_label('Tax ID', exact=True).fill('27AAAAA0000A1Z5')
        pp.get_by_role('button', name='Save as Preset', exact=True).click(); pp.get_by_label('Preset name').fill('My Shop'); pp.get_by_role('button', name='Save', exact=True).click()
        until(lambda: 'My Shop' in pp.evaluate("[...document.querySelectorAll('#root .pbar option')].map(o => o.textContent).join()"), 5000)
        pp.evaluate("localStorage.removeItem('tools.invoice.v1')"); pp.reload(); until(lambda: pp.evaluate("!!document.querySelector('#root .pbar select:not([disabled])')"), 8000)
        empty = pp.get_by_label('Business Name', exact=True).input_value(); pp.select_option('#root .pbar select', label='My Shop')
        until(lambda: pp.get_by_label('Business Name', exact=True).input_value() == 'Sharma Traders', 3000)
        check('R5 Invoice Generator: your business details come back from a preset', empty == '' and pp.get_by_label('Tax ID', exact=True).input_value() == '27AAAAA0000A1Z5', empty)
        pp.goto(U + '#/settings'); until(lambda: pp.evaluate("document.querySelectorAll('#root .pgroup').length") == 3, 8000)
        groups = pp.evaluate("[...document.querySelectorAll('#root .pgroup h3')].map(x => x.textContent)")
        check('R6 Tools Settings > Presets lists every preset by tool', groups == ['Sign PDF: Signatures', 'QR Code Styles', 'Invoice: Your Business'], groups)
        pp.get_by_role('button', name='Rename Blue Code', exact=True).click(); pp.get_by_label('New name for Blue Code').fill('Navy Code'); pp.keyboard.press('Enter')
        until(lambda: 'Navy Code' in pp.inner_text('#root .view'), 5000)
        pp.get_by_role('button', name='Delete My Signature', exact=True).click(); until(lambda: pp.evaluate("document.querySelectorAll('#root .pgroup').length") == 2, 5000)
        k = s.get('presets').get('kinds') or {}
        check('R7 ...rename and delete work, on the server', [p['name'] for p in k.get('qr', [])] == ['Navy Code'] and 'sign' not in k, k.keys())
        ov = s.get('backup_overview')['parts'].get('tools', {})
        check('R8 the site backup includes Tools presets', ov.get('summary') == '2 saved presets', ov)
        # 1.61.5: a choice of two or three is a row of buttons; on/off is a switch
        pp.goto(U + '#/tool/id-photo'); until(lambda: pp.evaluate("!!document.querySelector('#root .seg[role=radiogroup]')"), 8000)
        sg = pp.evaluate("(() => { const g = document.querySelector('#root .seg[role=radiogroup]'), sel = g.previousElementSibling; return [sel && sel.tagName, sel && getComputedStyle(sel).display, g.children.length, [...g.children].filter(b => b.getAttribute('aria-checked') === 'true').length]; })()")
        check('R10 a tool\'s choice of two or three is a row of buttons (the select stays, hidden, behind it)', sg[0] == 'SELECT' and sg[1] == 'none' and sg[2] in (2, 3) and sg[3] == 1, sg)
        before = pp.evaluate("document.querySelector('#root .seg[role=radiogroup]').previousElementSibling.value")
        pp.evaluate("[...document.querySelectorAll('#root .seg[role=radiogroup] button')].find(b => b.getAttribute('aria-checked') !== 'true').click()")
        check('R11 pressing another button changes the choice the tool reads', pp.evaluate("document.querySelector('#root .seg[role=radiogroup]').previousElementSibling.value") != before)
        until(lambda: pp.evaluate("!!document.querySelector('#root .chk input[type=checkbox]')"), 8000)
        sw = pp.evaluate("(() => { const c = document.querySelector('#root .chk input[type=checkbox]'); return [Math.round(c.getBoundingClientRect().width), getComputedStyle(c).appearance]; })()")
        check('R12 an on/off option is a switch', sw[0] == 42 and sw[1] == 'none', sw)
        check('R9 no page error with presets', not perr, perr)
        pc.close()
        b.close()
finally:
    stop_all()
sys.exit(check.done())
