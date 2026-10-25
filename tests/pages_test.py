#!/usr/bin/env python3
"""Every page and every Settings section in a real browser, with demo data: no page error, no broken request, and the Cards Excel export opens.
  python3 tests/pages_test.py build/site-update-X.Y.Z.zip
Needs PHP 8.3+ and Playwright with Chromium (or --browser firefox|webkit); openpyxl is used to read the Excel file when it is installed.
Nothing leaves this computer: every request to another site is blocked."""
import io, json, os, re, shutil, subprocess, sys, tempfile, time, urllib.request, http.cookiejar, zipfile
from playwright.sync_api import sync_playwright

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from helpers import browser_name, launch, browser_login, CSP_REPORTER
BROWSER = browser_name()
if len(sys.argv) < 2: sys.exit(__doc__)
ZIP = sys.argv[1]; PORT = 8195; B = 'http://127.0.0.1:%d' % PORT; UA = 'PagesTest/1.0'; PW = 'correct horse battery staple'
D = tempfile.mkdtemp(); zipfile.ZipFile(ZIP).extractall(D)
php = subprocess.Popen(['php', '-S', '127.0.0.1:%d' % PORT, '-t', D], env=dict(os.environ, PHP_CLI_SERVER_WORKERS='6'), stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
time.sleep(1)
res = []
def check(name, ok, extra=''):
    res.append(bool(ok)); print(('PASS ' if ok else 'FAIL ') + name + ('' if ok or not extra else '  -> ' + str(extra)[:300]))

cj = http.cookiejar.CookieJar(); op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj))
def call(a, b=None):
    rq = urllib.request.Request(B + '/web/core-data/api.php?action=' + a, data=json.dumps(b).encode() if b is not None else None, headers={'User-Agent': UA, 'Content-Type': 'application/json'}, method='POST' if b is not None else 'GET')
    return json.loads(op.open(rq).read())
try:
    call('state')
    code = re.search(r'([A-Z0-9]{4}(?:-[A-Z0-9]{4})+|[A-Z0-9]{8,})', open(D + '/site-data/core/SETUP-CODE.php').read().replace('<?php', '')).group(1)
    check('P1 a new site can be set up', call('setup', {'code': code, 'password': PW}).get('ok'))
    subprocess.run(['php', os.path.join(HERE, 'seed_demo.php'), D], capture_output=True)
    ver = re.search(r"define\('HUB', '([^']+)'\)", open(D + '/web/core-data/lib/boot.php').read()).group(1)
    cookies = [{'name': c.name, 'value': c.value, 'url': B + '/'} for c in cj]
    panes = re.findall(r"\['([a-z]+)', '[A-Z][^']*', ICON", open(D + '/web/docs/src/core-data/shell-admin.src.js', encoding='utf-8').read())
    with sync_playwright() as p:
        b = launch(p, BROWSER); print('browser: %s %s' % (BROWSER, b.version))
        ctx = b.new_context(user_agent=UA, accept_downloads=True, viewport={'width': 1280, 'height': 860})
        ctx.add_init_script("try{localStorage.setItem('ix-seen-ver','%s')}catch(e){}" % ver)
        ctx.add_init_script(CSP_REPORTER)
        # PHP's test server ignores .htaccess, so each page gets the security policy Apache would send, read from its own .htaccess
        CSP = {}
        for app in ('links', 'info', 'cards', 'tasks', 'tools', 'settings'):
            m = re.search(r'Header (?:always )?set Content-Security-Policy "([^"]+)"', open(D + '/web/%s/.htaccess' % app).read()) if os.path.isfile(D + '/web/%s/.htaccess' % app) else None
            if m: CSP['/web/%s/' % app] = m.group(1)
        def handle(r):
            u = r.request.url
            if not (u.startswith(B) or u.startswith(('data:', 'blob:'))): return r.abort()
            path = u[len(B):].split('#')[0].split('?')[0]
            # WebKit fails the requests of a page that Playwright fetched and handed over (seen in CI: P6 failed, P7 passed),
            # so there each page keeps only its own <meta> policy; Chromium, Firefox and the Apache job check the headers.
            if BROWSER != 'webkit' and r.request.resource_type == 'document' and path in CSP:
                resp = r.fetch(); h = dict(resp.headers); h['content-security-policy'] = CSP[path]; return r.fulfill(response=resp, headers=h)
            r.continue_()
        ctx.route('**/*', handle)
        pg = ctx.new_page(); errs = []; bad = []
        why = browser_login(pg, B, PW)
        check('P0 the browser logs in through the login dialog', why == '', why)
        pg.on('pageerror', lambda e: errs.append(str(e)[:200]))
        refused = []; tools_main = set()
        pg.on('console', lambda m: refused.append(m.text[:160]) if m.type in ('error', 'warning') and ('Content Security Policy' in m.text or 'Content-Security-Policy' in m.text or 'Refused to' in m.text or 'React error' in m.text) else None)
        pg.on('request', lambda r: tools_main.add(r.url.replace(B, '')) if '/tools/assets/app-' in r.url and r.url.split('?')[0].endswith('.js') else None)
        pg.on('response', lambda r: bad.append((r.status, r.url.replace(B, ''))) if r.url.startswith(B) and r.status >= 400 else None)
        problems = []; noicon = []
        for name, url in [('Home', '/'), ('Links', '/web/links/'), ('Info', '/web/info/'), ('Cards', '/web/cards/'), ('Tasks', '/web/tasks/'), ('Tools', '/web/tools/'), ('Tools: a tool', '/web/tools/#/tool/image-resize')] + [('Settings ' + s, '/web/settings/#' + s) for s in panes]:
            errs.clear(); bad.clear()
            pg.goto(B + url, wait_until='domcontentloaded'); pg.wait_for_timeout(1100)
            if errs or bad: problems.append((name, errs[:1], bad[:2]))
            ic = pg.evaluate("(document.querySelector('link[rel~=icon]') || {}).href || ''")
            if not ic.endswith(('/brand/favicon-64.png', '/brand/logo-192.png')) and '/media/' not in ic: noicon.append((name, ic[:60]))
        check('P2 every page and all %d Settings sections load with no page error and no broken request' % len(panes), len(panes) >= 8 and not problems, problems)
        check('P9 every page, Tools included, shows the site logo as its tab icon', not noicon, noicon)
        check('P4 with the real security policies, nothing is refused and React reports no error', len(CSP) >= 5 and not refused, (len(CSP), refused[:3]))
        check('P5 Tools loads its main file once', len(tools_main) == 1, sorted(tools_main))
        pg.goto(B + '/web/cards/'); pg.wait_for_timeout(1500)
        admin = pg.evaluate("!!(window.HUB && HUB.isAdmin && HUB.isAdmin())")
        diag = '' if admin else (
            [(c['name'], c['domain'], c['path'], c.get('sameSite'), c.get('secure'), c.get('expires')) for c in ctx.cookies()],
            pg.evaluate("fetch('../core-data/api.php?action=state', {credentials: 'same-origin'}).then(r => r.json()).then(j => ({admin: j.admin}), e => String(e))"))
        check('P6 the browser is logged in (the login cookie was accepted)', admin, diag)
        # P7: the same, in a context without the test's request routing (the routing hands each page to the browser with its
        # security policy added). If P7 passes where P6 fails, the difference is in the routing, not in the site.
        c7 = b.new_context(user_agent=UA, viewport={'width': 1280, 'height': 860}); p7 = c7.new_page()
        why7 = browser_login(p7, B, PW); p7.goto(B + '/web/cards/'); p7.wait_for_timeout(1500)
        check('P7 logged in on a second page without the test routing', p7.evaluate("!!(window.HUB && HUB.isAdmin && HUB.isAdmin())"), why7)
        c7.close()
        # P8: Text Sizes now opens from each page's own settings dialog, and is no longer its own item in the pencil menu
        p8 = []
        for app, dlg, btn, done in (('links', 'setDlg', 'sTextBtn', "document.getElementById('textDlg').open"),
                                    ('info', 'setdlg', 'setsizes', "!!document.querySelector('#ix-dlg.open #ix-tzs')"),
                                    ('cards', 'tdlg', 'tsizes', "!!document.querySelector('#ix-dlg.open #ix-tzs')")):
            pg.goto(B + '/web/%s/' % app, wait_until='domcontentloaded'); pg.wait_for_timeout(1200)
            item = pg.evaluate("document.querySelectorAll('[data-hub=item][data-hub-label=\"Text Sizes\"]').length")
            pg.evaluate("document.getElementById('%s').showModal(); document.getElementById('%s').click()" % (dlg, btn)); pg.wait_for_timeout(700)
            ok = pg.evaluate(done) and not pg.evaluate("document.getElementById('%s').open" % dlg)
            if item or not ok: p8.append((app, item, ok))
        check('P8 Text Sizes opens from the Settings dialog of Links, Info and Cards (and is not a separate pencil-menu item)', not p8, p8)
        data = b''
        try:
            with pg.expect_download(timeout=20000) as dl: pg.evaluate("document.querySelector('[data-a=xlsx]').click()")
            data = open(dl.value.path(), 'rb').read()
        except Exception as e:
            check('P3 Cards exports a file (no download came)', False, str(e)[:120])
        try:
            import openpyxl
            if not data: raise ValueError('no file')
            wb = openpyxl.load_workbook(io.BytesIO(data)); ws = wb['Cards']; hdr = [c.value for c in ws[1]]
            ok = 'Limit' in hdr and isinstance(ws.cell(2, hdr.index('Limit') + 1).value, (int, float)) and '₹' in ws.cell(2, hdr.index('Limit') + 1).number_format
            detail = (wb.sheetnames, hdr[:5])
        except ValueError:
            ok = False; detail = 'no file'
        except ImportError:
            z = zipfile.ZipFile(io.BytesIO(data)); ok = 'xl/worksheets/sheet1.xml' in z.namelist(); detail = z.namelist()
        check('P3 Cards exports a real Excel file with numbers and the money format', data[:2] == b'PK' and ok, detail)
        # P10-P16: the 1.48.5 Cards page (Table or By Holder sections, four total boxes, holder chips) and the tidied pencil menus
        def items(app):
            pg.goto(B + '/web/%s/' % app, wait_until='domcontentloaded'); pg.wait_for_timeout(900)
            return pg.evaluate("[...document.querySelectorAll('[data-hub=item]')].map((e, i) => [+(e.dataset.hubOrder || 0), i, e.dataset.hubLabel]).sort((x, y) => x[0] - y[0] || x[1] - y[1]).map(x => x[2])")
        want = {'links': ['Colors', 'Fetch Icons', 'Import Bookmarks', 'Links Settings'], 'info': ['Export', 'Info Settings'],
                'tasks': ['Add Task', 'Spreadsheet (CSV)', 'Calendar File (.ics)', 'Print or PDF', 'Tasks Settings'],
                'cards': ['Add Card', 'Add Holder', 'Edit Holders', 'Edit Banks', 'PDF (print)', 'Excel Spreadsheet', 'Cards Settings']}
        got = {k: items(k) for k in want}
        check('P10 the pencil menus: Links (Colors, Fetch Icons, Import Bookmarks, Links Settings), Info, Tasks and Cards name their own settings', got == want, got)
        pg.goto(B + '/web/cards/', wait_until='domcontentloaded'); pg.evaluate("localStorage.removeItem('cards-view'); localStorage.removeItem('cards-fold')")
        pg.reload(wait_until='domcontentloaded'); pg.wait_for_selector('#list .row.item', timeout=15000)
        grp = lambda: pg.evaluate("[...document.querySelectorAll('#list .grp .gn')].map(e => e.textContent)")
        rows = lambda: pg.evaluate("document.querySelectorAll('#list .row.item').length")
        stat = lambda: pg.evaluate("[...document.querySelectorAll('#stats .stat')].map(e => e.querySelector('.sl').textContent + '=' + e.querySelector('b').textContent)")
        first = (len(grp()), rows(), pg.evaluate("getComputedStyle(document.querySelector('#list .row.item .co')).display"), stat())
        check('P11 Cards opens as the Table (every card, with a Holder column) under four boxes: Cards, Lifetime Free, Charges a Year, Expiring Soon; the first chip is All Cards',
              first[:3] == (0, 9, 'block') and first[3] == ['Cards=9', 'Lifetime Free=4 of 9', 'Charges a Year=₹8,499', 'Expiring Soon=1'] and pg.evaluate("document.querySelector('#fchips [data-h=All]').textContent") == 'All Cards'
              and pg.evaluate("document.querySelector('#fview [aria-pressed=true]').textContent") == 'Table', first)
        pg.click('#fview [data-v=holder]'); pg.wait_for_timeout(200)
        check('P12 By Holder: one section per holder with its totals, and the choice is kept',
              grp() == ['Alex', 'Sam', 'Jordan'] and rows() == 9 and pg.evaluate("document.querySelector('#list .grp .gh').textContent.includes('4 cards')") and pg.evaluate("localStorage.getItem('cards-view')") == 'holder', grp())
        pg.click('#list .gh[data-fold="Jordan"]'); pg.wait_for_timeout(200); folded = rows()
        pg.reload(wait_until='domcontentloaded'); pg.wait_for_selector('#list .row.item', timeout=15000)
        check('P13 a holder section folds, and stays folded (and By Holder stays on) after a reload', folded == 7 and rows() == 7 and len(grp()) == 3, (folded, rows()))
        pg.click('#list .gh[data-fold="Jordan"]'); pg.click('#fview [data-v=list]'); pg.wait_for_timeout(200)
        pg.click('#fchips [data-h="Sam"]'); pg.wait_for_timeout(200)
        sam = (rows(), stat()[0], pg.evaluate("!document.getElementById('fclear').hidden"))
        pg.click('#fclear'); pg.wait_for_timeout(200)
        check('P16 a holder chip filters to that holder (the boxes follow), and Clear brings every card back', sam == (3, 'Cards=3', True) and rows() == 9, sam)
        pg.click('#list .row.item[data-id] .cn small'); pg.wait_for_timeout(300)
        opened = pg.evaluate("document.getElementById('dlg').open"); pg.evaluate("document.getElementById('dlg').close()")
        pg.evaluate("document.getElementById('tdlg').showModal()"); pg.click('#tcols'); pg.wait_for_timeout(400)
        cols = pg.evaluate("document.getElementById('odlg').open && document.querySelectorAll('#olist .orow').length"); pg.evaluate("document.getElementById('odlg').close()")
        check('P14 clicking a card row opens Edit Card, and Cards Settings opens Column Order', opened and cols == 11, (opened, cols))
        pctx = b.new_context(user_agent=UA, viewport={'width': 390, 'height': 844}); pctx.add_cookies(cookies); pctx.add_init_script("try{localStorage.setItem('ix-seen-ver','%s')}catch(e){}" % ver)
        pctx.route('**/*', handle); pp = pctx.new_page(); pp.goto(B + '/web/cards/', wait_until='domcontentloaded'); pp.wait_for_selector('#list .row.item', timeout=15000)
        ph = pp.evaluate("({sw: document.documentElement.scrollWidth - innerWidth, head: getComputedStyle(document.querySelector('#list .row.head')).display, tall: Math.max(...[...document.querySelectorAll('#list .row.item')].map(r => r.getBoundingClientRect().height))})")
        check('P15 Cards on a phone: nothing scrolls sideways, no column headings, and each card takes two short lines', ph['sw'] <= 1 and ph['head'] == 'none' and ph['tall'] < 80, ph)
        pctx.close()
        b.close()
finally:
    php.terminate(); shutil.rmtree(D, ignore_errors=True)
print('\n%d/%d passed' % (sum(res), len(res)))
sys.exit(0 if all(res) else 1)
